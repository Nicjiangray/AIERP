"""AI·ERP 后端：数据库表 API + 与前端约定的 /ai/chat（直连 LLM）。"""

from __future__ import annotations

from dotenv import load_dotenv

load_dotenv()

import asyncio
import asyncpg
import io
import json
import logging

import bcrypt
import os
import re
import sys
import uuid
from contextlib import asynccontextmanager
from datetime import date, datetime, time, timezone
from decimal import Decimal
from pathlib import Path
from typing import Any, Dict, List, Optional

import psycopg2.extras
from fastapi import Body, FastAPI, HTTPException, Query, Request
from fastapi.encoders import jsonable_encoder
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import Response, StreamingResponse
from pydantic import BaseModel, ConfigDict, Field
from psycopg2.extras import Json

from tools.database import close_all_pools, db_conn
from engine.memory_manager import (
    clean_session,
    compress_longterm,
    extract_core_facts,
    maybe_clean_previous_session,
)
from engine.incident_logger import log_incident
from engine.prompt_builder import build_prompt
from engine.sql_engine import (
    confirm_batch_sql,
    execute_sql,
    highlight_words_from_summary,
    preview_sql,
    user_can_read_table,
    verify_confirmed_writes,
)
from llm.client import (
    chat,
    chat_expert_only,
    chat_stream,
    parse_file,
)
from engine.task_manager import (
    advance_step,
    create_task,
    get_active_task,
    get_task_progress_text,
    is_last_step,
    parse_plan,
    update_step_result,
)

logger = logging.getLogger(__name__)
logger.setLevel(logging.INFO)
if not logger.handlers:
    _orch_h = logging.StreamHandler(sys.stderr)
    _orch_h.setLevel(logging.INFO)
    _orch_h.setFormatter(logging.Formatter("%(levelname)s:%(name)s:%(message)s"))
    logger.addHandler(_orch_h)
logger.propagate = False

# asyncpg 池：供 engine.task_manager（business.ai_tasks）使用，与 psycopg2 的 db_conn 并行
_task_pg_pool: Optional[asyncpg.Pool] = None

# --- 外层 orchestrator（与内层 expert 步数无关，勿混用）---
# 主助手每轮回复 ↔ 解析 [TOOL] ↔ 调一次 _call_tool_expert（内层另有 MAX_EXPERT_STEPS）的上限。
MAX_ORCHESTRATOR_ROUNDS = 30
# 历史名称：与 MAX_ORCHESTRATOR_ROUNDS 同一概念，非 expert 内层步数。
MAX_TOOL_ROUNDS = MAX_ORCHESTRATOR_ROUNDS

_ORCH_MSG_MAX_ROUNDS = (
    "任务步骤过多，已暂停。你可以继续对话推进剩余步骤。"
)
_ORCH_MSG_DUP_TOOL = "检测到重复操作，已暂停。请检查后继续。"
_ORCH_SYSTEM_ABORT = "用户已终止自动执行。"

# --- 内层 tool_expert（_run_tool_expert_dialog）---
# 单次 [TOOL] 调用里 query/write/done 的多轮对话上限；与外层 orchestrator 轮数独立。
MAX_EXPERT_STEPS = 30
# 用户确认写入后：执行或校验失败时专家重写 SQL 重试，与上面「专家总步数」无关
MAX_WRITE_CONFIRM_ATTEMPTS = 3


def _extract_json_from_text(text: str) -> Optional[Dict[str, Any]]:
    """从文本中提取第一个 JSON 对象（支持 ```json 代码块；正文用 JSONDecoder.raw_decode 处理嵌套）。"""
    if not text or not isinstance(text, str):
        return None
    raw = text.strip()
    for m in re.finditer(r"```(?:json)?\s*([\s\S]*?)```", raw, re.I):
        chunk = m.group(1).strip()
        try:
            obj = json.loads(chunk)
            if isinstance(obj, dict):
                return obj
        except json.JSONDecodeError:
            continue
    dec = json.JSONDecoder()
    start = raw.find("{")
    while start >= 0:
        try:
            obj, _end = dec.raw_decode(raw, start)
            if isinstance(obj, dict):
                return obj
        except json.JSONDecodeError:
            pass
        start = raw.find("{", start + 1)
    return None


def _extract_tool_block(text: str) -> Optional[Dict[str, str]]:
    """从 LLM 回复中提取 [TOOL]...[/TOOL] 块，解析为 dict。"""
    m = re.search(r"\[TOOL\]\s*(.*?)\s*\[/TOOL\]", text, re.DOTALL | re.I)
    if not m:
        return None
    block: Dict[str, str] = {}
    for line in m.group(1).strip().splitlines():
        line = line.strip()
        if not line or ":" not in line:
            continue
        key, _, val = line.partition(":")
        block[key.strip()] = val.strip()
    if "action" not in block:
        return None
    logger.info(
        f"orchestrator: 检测到 [TOOL] 块: action={block.get('action')}, intent={block.get('intent')}"
    )
    return block


def _strip_plan_for_stream_visible(full_reply: str) -> str:
    """
    从累积正文中去掉 [PLAN]...[/PLAN]（闭合则删块；仅出现 [PLAN] 未闭合则截断到其前），
    供 SSE 与最终可见文本一致。
    """
    if not full_reply:
        return full_reply
    text = full_reply
    while True:
        pi = text.find("[PLAN]")
        if pi < 0:
            break
        end = text.find("[/PLAN]", pi)
        if end < 0:
            return text[:pi].rstrip()
        text = text[:pi] + text[end + len("[/PLAN]") :]
    return text


def _user_visible_stream_prefix(full_reply: str) -> str:
    """
    SSE 发给用户可见的文本：不包含 [PLAN]...[/PLAN]；不包含 [TOOL] 块内原文，也不包含 [/TOOL] 之后同轮编造内容。
    若已出现 [TOOL] 但未闭合，只展示 [TOOL] 之前的部分（块内缓冲不下发）。
    若流式片段中仅出现不完整的「[TOOL」（尚未写出闭合 ]），也截断在该标记之前，避免把「[TOOL」字样下发给用户。
    """
    if not full_reply:
        return full_reply
    full_reply = _strip_plan_for_stream_visible(full_reply)
    if not full_reply:
        return full_reply
    low = full_reply.lower()
    ti = low.find("[tool]")
    if ti >= 0:
        end = low.find("[/tool]", ti)
        if end < 0:
            return full_reply[:ti]
        return full_reply[:ti]
    # 仅有未写完的 [TOOL 开头（尚无 ]），整段不展示从该处起的缓冲
    partial = re.search(r"\[tool(?!])", low)
    if partial:
        return full_reply[: partial.start()]
    return full_reply


def _strip_assistant_fabrication_after_tool(full_reply: str) -> str:
    """
    写入 ai_chat、并作为下一轮 assistant 消息时，去掉 [/TOOL] 之后的同轮续写，避免把编造带进上下文。
    """
    low = full_reply.lower()
    ti = low.find("[tool]")
    if ti < 0:
        return full_reply
    end = low.find("[/tool]", ti)
    if end < 0:
        return full_reply
    return full_reply[: end + len("[/TOOL]")].rstrip()


def _expert_exec_feedback(exec_result: Dict[str, Any]) -> str:
    """给 tool_expert 的反馈：失败时把 error 放在最前，便于模型纠错。"""
    body = json.dumps(exec_result, ensure_ascii=False, default=str)
    if not exec_result.get("success"):
        err = exec_result.get("error") or "未知错误"
        return f"执行失败。\n首要错误：{err}\n完整结果（JSON）：\n{body}"
    return f"执行结果：\n{body}"


_EXPERT_LOG_SQL_MAX = 2500
_EXPERT_LOG_FEEDBACK_MAX = 4000
_EXPERT_LOG_REPLY_HEAD = 600

# 流式对话：工具执行前提示；写入预览前提示（并记入 ai_chat）
_SSE_TOOL_PLEASE_WAIT = "\n\n请稍等......\n\n"
_WRITE_PREVIEW_CONFIRM_HINT = "「确认的内容」请确认。"


def _write_preview_user_visible_log(result: Dict[str, Any]) -> str:
    """与前端写入预览卡片一致的可读摘要，写入 business.ai_chat 助手行，便于用户与清洗任务可见。"""
    lines = [_WRITE_PREVIEW_CONFIRM_HINT]
    summary = (result.get("summary") or "").strip()
    if summary:
        lines.extend(["", "操作确认", summary])
    exp = (result.get("expert_explanation") or "").strip()
    if exp:
        lines.extend(["", "操作说明", exp])
    return "\n".join(lines).strip()


def _trunc_for_log(s: Optional[str], max_len: int = _EXPERT_LOG_SQL_MAX) -> str:
    if s is None:
        return ""
    t = str(s)
    if len(t) <= max_len:
        return t
    return t[:max_len] + f"...(truncated total_len={len(t)})"


def _format_tool_result(tool_block: Dict[str, str], result: Dict[str, Any]) -> str:
    """把 tool_expert 最终结果格式化为 [TOOL_RESULT] 块注入给助手。"""
    action = tool_block.get("action", "")
    lines = ["[TOOL_RESULT]", f"action: {action}"]

    if not result.get("success"):
        lines.append("status: error")
        lines.append(f"error: {result.get('error', '执行失败')}")
        if result.get("need_info"):
            lines.append(f"need_info: {result['need_info']}")
        if result.get("candidates") is not None:
            c = result["candidates"]
            lines.append(
                f"candidates: {c if isinstance(c, str) else json.dumps(c, ensure_ascii=False, default=str)}"
            )
    elif result.get("stage") == "preview" and (
        result.get("sql") or result.get("batch_writes")
    ):
        lines.append("status: pending_confirm")
        lines.append(f"summary: {result.get('summary', '')}")
        bw = result.get("batch_writes")
        if bw:
            lines.append(f"batch_count: {len(bw)}")
    else:
        lines.append("status: success")
        lines.append(f"summary: {result.get('summary', '')}")

    lines.append("[/TOOL_RESULT]")
    result_text = "\n".join(lines)
    logger.info(f"orchestrator: 回注给助手的 TOOL_RESULT:\n{result_text}")
    return result_text


def _append_progress_line_to_tool_result(text: str, progress: str) -> str:
    progress = (progress or "").strip()
    if not progress or "[/TOOL_RESULT]" not in text:
        return text
    return text.replace("[/TOOL_RESULT]", f"{progress}\n[/TOOL_RESULT]", 1)


def _expert_result_is_done_success(result: Dict[str, Any]) -> bool:
    """工具专家以 type=done 结束且 success=True（与 _run_tool_expert_dialog 返回字段一致）。"""
    if not result.get("success"):
        return False
    return str(result.get("type") or "").strip().lower() == "done"


async def _format_tool_result_with_task_progress(
    tool_block: Dict[str, str],
    result: Dict[str, Any],
    session_id: Optional[int],
    user_id: Optional[int],
) -> str:
    """
    在工具专家返回的 TOOL_RESULT 基础上，当专家以 done+success 结束时补充任务进度行（见 get_task_progress_text）。
    """
    base = _format_tool_result(tool_block, result)
    if not _expert_result_is_done_success(result):
        return base
    if session_id is None or user_id is None or _task_pg_pool is None:
        return base
    progress = await get_task_progress_text(_task_pg_pool, session_id, user_id)
    if not progress:
        return base
    return _append_progress_line_to_tool_result(base, progress)


def _expert_payload_is_done(payload: Dict[str, Any]) -> bool:
    """与 tool_expert 约定：done: true 或 type: done 均表示结束本轮（含 success: false 的说明性结束）。"""
    if str(payload.get("type") or "").strip().lower() == "done":
        return True
    if payload.get("done"):
        return True
    return False


def _combine_expert_done_failure_message(payload: Dict[str, Any]) -> str:
    """将 done 且失败时的 error / need_info 原样拼成一段，便于注入 [TOOL_RESULT] 与助手理解。"""
    parts: List[str] = []
    err = payload.get("error")
    if err is not None and str(err).strip() != "":
        parts.append(str(err).strip())
    ni = payload.get("need_info")
    if ni is not None and str(ni).strip() != "":
        parts.append(str(ni).strip())
    return "\n".join(parts) if parts else "执行未成功"


def _is_write_preview_result(result: Dict[str, Any]) -> bool:
    """tool_expert 返回的 preview_sql 成功结果（待用户确认）。"""
    if not result.get("success"):
        return False
    if result.get("stage") != "preview":
        return False
    if result.get("batch_writes"):
        return len(result["batch_writes"]) > 0
    if not result.get("sql"):
        return False
    op = (result.get("operation") or "").lower()
    return op in ("insert", "update", "delete")


async def _run_tool_expert_dialog(
    expert_messages: List[Dict[str, Any]],
    user_id: Optional[int] = None,
    session_id: Optional[int] = None,
    trace_user_input: Optional[str] = None,
) -> Dict[str, Any]:
    """tool_expert 多轮对话：query / write / done，与入口消息无关（由调用方注入）。"""
    pending_writes: List[Dict[str, Any]] = []

    def _run_expert() -> str:
        return chat(expert_messages)

    for step in range(MAX_EXPERT_STEPS):
        logger.info(
            "tool_expert_detail | step_begin | n=%d/%d | dialog_messages=%d",
            step + 1,
            MAX_EXPERT_STEPS,
            len(expert_messages),
        )
        logger.info(f"tool_expert: 第 {step + 1}/{MAX_EXPERT_STEPS} 步")
        try:
            expert_reply = await asyncio.to_thread(_run_expert)
        except Exception as e:
            log_incident(
                stage="llm_call",
                error_message=str(e),
                session_id=session_id,
                user_id=user_id,
                user_input=trace_user_input,
                context={"step": step + 1, "dialog": "tool_expert"},
            )
            return {"success": False, "error": str(e)}
        logger.info(f"tool_expert: 回复:\n{expert_reply}")
        logger.info(
            "tool_expert_detail | model_raw | n=%d | chars=%d | head=%r",
            step + 1,
            len(expert_reply or ""),
            _trunc_for_log(expert_reply, _EXPERT_LOG_REPLY_HEAD),
        )

        payload = _extract_json_from_text(expert_reply)
        if not payload:
            log_incident(
                stage="parse",
                error_message="工具专家未返回有效 JSON 指令",
                session_id=session_id,
                user_id=user_id,
                user_input=trace_user_input,
                context={
                    "raw_head": (expert_reply or "")[:2000],
                    "step": step + 1,
                },
            )
            return {"success": False, "error": "工具专家未返回有效指令"}

        logger.info(f"orchestrator: expert payload: {payload}")
        logger.info(
            "tool_expert_detail | parsed | n=%d | done=%s | type=%s | sql=%s | params=%s | explanation_head=%r",
            step + 1,
            payload.get("done"),
            payload.get("type"),
            _trunc_for_log((payload.get("sql") or "").strip() or None),
            payload.get("params"),
            _trunc_for_log(str(payload.get("explanation") or ""), 200),
        )

        if _expert_payload_is_done(payload):
            if pending_writes:
                if not payload.get("success", False):
                    msg = _combine_expert_done_failure_message(payload)
                    if msg == "执行未成功":
                        msg = (
                            payload.get("error")
                            or "存在未确认的写入预览，但 done 未标记成功"
                        )
                    return {"success": False, "error": msg}
                expert_overall = str(payload.get("summary") or "").strip() or None
                engine_lines = [
                    f"{i + 1}) {w.get('summary') or ''}"
                    for i, w in enumerate(pending_writes)
                ]
                overall_engine = "\n".join(engine_lines).strip() or "请确认以下写入操作"
                batch_out = {
                    "success": True,
                    "stage": "preview",
                    "batch_writes": pending_writes,
                    "summary": overall_engine,
                    "expert_explanation": expert_overall,
                    "highlight_words": highlight_words_from_summary(overall_engine),
                    "payload": {
                        "batch": [
                            {
                                "sql": w["sql"],
                                "params": w.get("params") or [],
                                "table": w.get("table"),
                                "operation": w.get("operation"),
                                "expert_explanation": w.get("expert_explanation"),
                            }
                            for w in pending_writes
                        ],
                        "summary": overall_engine,
                        "expert_explanation": expert_overall,
                    },
                    "operation": "batch",
                    "sql": pending_writes[0]["sql"],
                }
                logger.info(
                    "tool_expert_detail | done_batch_preview | n=%d | writes=%d | summary_head=%r",
                    step + 1,
                    len(pending_writes),
                    _trunc_for_log(overall_engine, 400),
                )
                return batch_out
            success_ok = bool(payload.get("success", False))
            summary = str(payload.get("summary", "") or "")
            if not success_ok:
                combined_err = _combine_expert_done_failure_message(payload)
                out = {
                    "success": False,
                    "summary": summary,
                    "error": combined_err,
                    "need_info": None,
                    "candidates": payload.get("candidates"),
                }
            else:
                out = {
                    "success": True,
                    "type": "done",
                    "summary": summary,
                    "error": payload.get("error"),
                    "need_info": payload.get("need_info"),
                    "candidates": payload.get("candidates"),
                }
            logger.info(
                "tool_expert_detail | done | n=%d | success=%s | summary=%r | error=%r | need_info=%r",
                step + 1,
                out.get("success"),
                _trunc_for_log(str(out.get("summary") or ""), 500),
                out.get("error"),
                payload.get("need_info"),
            )
            logger.info(
                f"orchestrator: engine 执行结果: success={out.get('success')}, keys={list(out.keys())}"
            )
            return out

        sql = (payload.get("sql") or "").strip()
        cmd_type = (payload.get("type") or "").strip().lower()
        raw_params = payload.get("params")
        params_list: List[Any] = raw_params if isinstance(raw_params, list) else []

        if cmd_type == "query":
            if not sql:
                return {"success": False, "error": "query 指令缺少 sql"}

            def _do_query() -> Dict[str, Any]:
                with db_conn() as conn:
                    return execute_sql(sql, params_list, conn, user_id=user_id)

            exec_result = await asyncio.to_thread(_do_query)
            logger.info(
                "tool_expert: query 结果: success=%s, count=%s, error=%r",
                exec_result.get("success"),
                exec_result.get("count"),
                exec_result.get("error"),
            )
            logger.info(
                "orchestrator: engine 执行结果: success=%s, keys=%s, error=%r",
                exec_result.get("success"),
                list(exec_result.keys()),
                exec_result.get("error"),
            )
            rows = exec_result.get("rows")
            row_preview = ""
            if isinstance(rows, list) and rows:
                try:
                    row_preview = _trunc_for_log(
                        json.dumps(rows[:3], ensure_ascii=False, default=str), 1200
                    )
                except Exception:
                    row_preview = "(rows preview failed)"
            logger.info(
                "tool_expert_detail | engine_query | n=%d | success=%s | count=%s | error=%r | rows_preview=%s",
                step + 1,
                exec_result.get("success"),
                exec_result.get("count"),
                exec_result.get("error"),
                row_preview or None,
            )
            fb = _expert_exec_feedback(exec_result)
            logger.info(
                "tool_expert_detail | feedback_to_expert | n=%d | chars=%d | body=%r",
                step + 1,
                len(fb),
                _trunc_for_log(fb, _EXPERT_LOG_FEEDBACK_MAX),
            )
            expert_messages.append({"role": "assistant", "content": expert_reply})
            expert_messages.append({"role": "user", "content": fb})
            continue

        if cmd_type == "write":
            if not sql:
                return {"success": False, "error": "write 指令缺少 sql"}
            explanation = str(payload.get("explanation") or "")

            def _do_preview() -> Dict[str, Any]:
                with db_conn() as conn:
                    return preview_sql(sql, params_list, conn, user_id=user_id)

            exec_result = await asyncio.to_thread(_do_preview)
            logger.info(
                "tool_expert: write preview 结果: success=%s, error=%r",
                exec_result.get("success"),
                exec_result.get("error"),
            )
            logger.info(
                "orchestrator: engine 执行结果: success=%s, keys=%s, error=%r",
                exec_result.get("success"),
                list(exec_result.keys()),
                exec_result.get("error"),
            )
            logger.info(
                "tool_expert_detail | engine_write_preview | n=%d | success=%s | table=%s | error=%r | sql=%s",
                step + 1,
                exec_result.get("success"),
                exec_result.get("table"),
                exec_result.get("error"),
                _trunc_for_log(sql),
            )
            if exec_result.get("success"):
                exp = explanation.strip()
                if exp:
                    exec_result["expert_explanation"] = exp
                pending_writes.append(
                    {
                        "sql": sql,
                        "params": params_list,
                        "table": exec_result.get("table"),
                        "operation": exec_result.get("operation"),
                        "summary": exec_result.get("summary"),
                        "expert_explanation": exp or None,
                    }
                )
                staged = len(pending_writes)
                # 仅一条写入时直接进入待确认（减少一轮「done」LLM）；多条时需再输出 done + 总述
                if staged == 1:
                    overall = pending_writes[0].get("summary") or "请确认以下写入"
                    overall_expert = pending_writes[0].get("expert_explanation") or None
                    logger.info(
                        "tool_expert_detail | auto_finalize_single_write | n=%d",
                        step + 1,
                    )
                    return {
                        "success": True,
                        "stage": "preview",
                        "batch_writes": list(pending_writes),
                        "summary": overall,
                        "expert_explanation": overall_expert,
                        "highlight_words": highlight_words_from_summary(overall),
                        "payload": {
                            "batch": [
                                {
                                    "sql": w["sql"],
                                    "params": w.get("params") or [],
                                    "table": w.get("table"),
                                    "operation": w.get("operation"),
                                    "expert_explanation": w.get("expert_explanation"),
                                }
                                for w in pending_writes
                            ],
                            "summary": overall,
                            "expert_explanation": overall_expert,
                        },
                        "operation": "batch",
                        "sql": pending_writes[0]["sql"],
                    }
                fb = (
                    f"已暂存第 {staged} 条写入预览（已通过审计，尚未执行）。\n"
                    f"{_expert_exec_feedback(exec_result)}\n\n"
                    "若还有写入，请继续输出一条 type 为 write 的 JSON；"
                    "若全部写入已规划完毕，请输出：\n"
                    '{"done": true, "success": true, "summary": "给用户看的整体说明（汇总将执行哪些写入）"}'
                )
                logger.info(
                    "tool_expert_detail | write_staged | n=%d | staged_total=%d",
                    step + 1,
                    staged,
                )
                expert_messages.append({"role": "assistant", "content": expert_reply})
                expert_messages.append({"role": "user", "content": fb})
                continue

            fb = _expert_exec_feedback(exec_result)
            logger.info(
                "tool_expert_detail | feedback_to_expert | n=%d | chars=%d | body=%r",
                step + 1,
                len(fb),
                _trunc_for_log(fb, _EXPERT_LOG_FEEDBACK_MAX),
            )
            expert_messages.append({"role": "assistant", "content": expert_reply})
            expert_messages.append({"role": "user", "content": fb})
            continue

        return {"success": False, "error": f"未知指令类型: {cmd_type}"}

    if pending_writes:
        _msg = f"已达最大步数（{MAX_EXPERT_STEPS}），仍有 {len(pending_writes)} 条写入待汇总；请输出 done 并附 summary，或拆分任务"
        log_incident(
            stage="expert",
            error_message=_msg,
            session_id=session_id,
            user_id=user_id,
            user_input=trace_user_input,
            context={
                "pending_staged": len(pending_writes),
                "max_steps": MAX_EXPERT_STEPS,
            },
        )
        return {
            "success": False,
            "error": _msg,
            "pending_staged": len(pending_writes),
        }
    log_incident(
        stage="expert",
        error_message="操作步骤过多，未能完成",
        session_id=session_id,
        user_id=user_id,
        user_input=trace_user_input,
        context={"max_steps": MAX_EXPERT_STEPS},
    )
    return {"success": False, "error": "操作步骤过多，未能完成"}


async def _call_tool_expert(
    tool_block: Dict[str, str],
    user_id: Optional[int] = None,
    session_id: Optional[int] = None,
) -> Dict[str, Any]:
    """
    把助手的意图指令发给 tool_expert；expert 内部多步执行（query / write / done），
    闭环完成后返回最终结果（摘要、错误或写预览）。
    """
    expert_system = build_prompt("tool_expert", user_id=user_id)
    user_msg = "\n".join(f"{k}: {v}" for k, v in tool_block.items())
    logger.info(f"orchestrator: 调用 tool_expert, 发送内容:\n{user_msg}")
    expert_messages: List[Dict[str, Any]] = [
        {"role": "system", "content": expert_system},
        {"role": "user", "content": user_msg},
    ]
    logger.info(
        "tool_expert_detail | start | max_steps=%d | intent_keys=%s",
        MAX_EXPERT_STEPS,
        list(tool_block.keys()),
    )
    return await _run_tool_expert_dialog(
        expert_messages,
        user_id=user_id,
        session_id=session_id,
        trace_user_input=user_msg,
    )


async def _call_tool_expert_recover_after_failed_confirm(
    failure_reason: str,
    failed_batch: List[Dict[str, Any]],
    last_confirm: Optional[Dict[str, Any]],
    tool_context: Optional[Dict[str, Any]] = None,
    user_id: Optional[int] = None,
    session_id: Optional[int] = None,
) -> Dict[str, Any]:
    """用户已点确认后执行或校验失败：由专家重写 batch（服务端自动再执行，无需用户再点预览）。"""
    expert_system = build_prompt("tool_expert", user_id=user_id)
    payload: Dict[str, Any] = {
        "phase": "recover_after_user_confirm",
        "说明": (
            "用户已确认过预览，但数据库执行或执行后校验未通过。"
            "请根据错误信息修正 SQL，重新输出 write（可多条），并以 done+summary 结束；"
            "仅一条 write 时可直接形成新预览。"
            "若「原始助手意图」中要求合同行项目/数量/产品，必须一并写入 contract_items，不可只写 contracts 主表。"
            "若 上次执行引擎返回 中 error_type 为 dependency_missing：请把 message 原样转告用户，"
            "并引导用户先在关联主数据（如员工、部门）中建立对应记录后再操作，勿盲目重试同一 SQL。"
        ),
        "错误与校验信息": failure_reason,
        "此前待执行的 batch": failed_batch,
    }
    if tool_context:
        payload["原始助手意图_TOOL块"] = tool_context
    if last_confirm is not None:
        payload["上次执行引擎返回"] = last_confirm
    user_msg = json.dumps(payload, ensure_ascii=False, default=str)
    expert_messages = [
        {"role": "system", "content": expert_system},
        {"role": "user", "content": user_msg},
    ]
    logger.info(
        "tool_expert_recover | start | batch_len=%d | reason_head=%r",
        len(failed_batch),
        _trunc_for_log(failure_reason, 300),
    )
    return await _run_tool_expert_dialog(
        expert_messages,
        user_id=user_id,
        session_id=session_id,
        trace_user_input=user_msg,
    )


def _insert_ai_chat_row(session_id: int, role: str, content: str) -> None:
    with db_conn() as conn:
        with conn.cursor() as cur:
            cur.execute(
                """
                INSERT INTO business.ai_chat (session_id, role, content)
                VALUES (%s, %s, %s)
                """,
                (session_id, role, content),
            )
        conn.commit()


def _env_bool(key: str, default: bool) -> bool:
    v = os.environ.get(key)
    if v is None:
        return default
    return v.strip().lower() in ("1", "true", "yes", "on")


def _cors_allow_origins_and_credentials() -> tuple[list[str], bool]:
    raw = (os.environ.get("CORS_ORIGINS") or "").strip()

    def _warn_star() -> None:
        if _env_bool("CORS_ALLOW_CREDENTIALS", False):
            logger.warning(
                "CORS_ALLOW_CREDENTIALS=true 与 CORS_ORIGINS=* 不兼容，已强制设为 false"
            )

    if raw == "*":
        _warn_star()
        return ["*"], False
    if not raw:
        origins = [
            "http://127.0.0.1:5173",
            "http://localhost:5173",
            "http://127.0.0.1:3000",
            "http://localhost:3000",
        ]
        return origins, _env_bool("CORS_ALLOW_CREDENTIALS", True)
    origins = [p.strip() for p in raw.split(",") if p.strip()]
    if not origins:
        return (
            ["http://127.0.0.1:5173", "http://localhost:5173"],
            _env_bool("CORS_ALLOW_CREDENTIALS", True),
        )
    if origins == ["*"]:
        _warn_star()
        return ["*"], False
    return origins, _env_bool("CORS_ALLOW_CREDENTIALS", True)


@asynccontextmanager
async def _app_lifespan(app: FastAPI):
    global _task_pg_pool
    _task_pg_pool = await asyncpg.create_pool(
        host=os.environ.get("PGHOST", "localhost"),
        database=os.environ.get("PGDATABASE", "aierp"),
        user=os.environ.get("PGUSER", "nicjiang"),
        password=os.environ.get("PG_API_PASSWORD", ""),
        port=int(os.environ.get("PGPORT", "5432")),
        min_size=1,
        max_size=10,
    )
    try:
        yield
    finally:
        if _task_pg_pool is not None:
            await _task_pg_pool.close()
            _task_pg_pool = None
        close_all_pools()


app = FastAPI(title="AI·ERP Backend", lifespan=_app_lifespan)
print(
    f"[aierp] 已加载 main.py: {Path(__file__).resolve()}",
    flush=True,
)

_cors_origins, _cors_cred = _cors_allow_origins_and_credentials()
print(
    f"[aierp] CORS allow_origins={_cors_origins!r} allow_credentials={_cors_cred}",
    flush=True,
)
app.add_middleware(
    CORSMiddleware,
    allow_origins=_cors_origins,
    allow_credentials=_cors_cred,
    allow_methods=["*"],
    allow_headers=["*"],
)

IDENT_RE = re.compile(r"^[a-zA-Z_][a-zA-Z0-9_]*$")
ALLOWED_SCHEMAS = frozenset({"meta", "business", "ai_assistant"})
BLOCK_COLUMNS = frozenset({"password_hash", "embedding", "intent_embedding"})


def _forbid_meta_table_api(schema: str) -> None:
    """通用表 CRUD/导出禁止走 HTTP 访问 meta（设计者请用 psql）。"""
    if (schema or "").strip().lower() == "meta":
        raise HTTPException(
            status_code=403,
            detail="系统表不可通过界面访问",
        )


def assert_ident(name: str, label: str = "identifier") -> None:
    if not name or not IDENT_RE.match(name):
        raise HTTPException(status_code=400, detail=f"Invalid {label}: {name}")


def fetch_column_meta(cur, schema: str, table: str) -> List[Dict[str, Any]]:
    cur.execute(
        """
        SELECT c.column_name AS name,
               c.data_type AS data_type,
               c.udt_name AS udt_name,
               c.is_nullable AS is_nullable,
               c.column_default AS column_default,
               c.is_generated AS is_generated,
               c.identity_generation AS identity_generation
        FROM information_schema.columns c
        WHERE c.table_schema = %s AND c.table_name = %s
        ORDER BY c.ordinal_position
        """,
        (schema, table),
    )
    return [dict(r) for r in cur.fetchall()]


def writable_columns(meta: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    out = []
    for r in meta:
        name = r["name"]
        if name in BLOCK_COLUMNS:
            continue
        if r.get("udt_name") == "vector":
            continue
        if (r.get("is_generated") or "NEVER") == "ALWAYS":
            continue
        out.append(r)
    return out


def coerce_cell(val: Any, udt_name: str, data_type: str, nullable: str) -> Any:
    if val is None:
        return None
    if val == "" and nullable == "YES":
        return None
    if isinstance(val, (dict, list)):
        return Json(val)
    return val


def table_has_id(meta: List[Dict[str, Any]]) -> bool:
    return any(r["name"] == "id" for r in meta)


LIKE_ESC = "!"


def filterable_column_names(meta: List[Dict[str, Any]]) -> List[str]:
    out = []
    for r in meta:
        n = r["name"]
        if n in BLOCK_COLUMNS:
            continue
        if r.get("udt_name") == "vector":
            continue
        out.append(n)
    return out


def _like_escape(s: str, esc: str = LIKE_ESC) -> str:
    if len(esc) != 1:
        raise ValueError("escape must be one character")
    return (
        s.replace(esc, esc + esc).replace("%", esc + "%").replace("_", esc + "_")
    )


def parse_filters_query(filters_raw: Optional[str]) -> Dict[str, str]:
    if not filters_raw or not str(filters_raw).strip():
        return {}
    try:
        o = json.loads(filters_raw)
    except json.JSONDecodeError:
        raise HTTPException(status_code=400, detail="filters 须为合法 JSON 对象")
    if not isinstance(o, dict):
        raise HTTPException(status_code=400, detail="filters 须为 JSON 对象")
    out: Dict[str, str] = {}
    for k, v in o.items():
        if not isinstance(k, str):
            continue
        vs = str(v).strip() if v is not None else ""
        if vs:
            out[k] = vs
    return out


def build_fuzzy_where_clause(
    meta: List[Dict[str, Any]],
    filters: Dict[str, str],
    q_global: Optional[str],
) -> tuple[str, List[Any]]:
    allowed = set(filterable_column_names(meta))
    parts: List[str] = []
    params: List[Any] = []
    esc = LIKE_ESC

    for col, val in filters.items():
        if col not in allowed or not IDENT_RE.match(col):
            continue
        pat = "%" + _like_escape(val, esc) + "%"
        parts.append(f'"{col}"::text ILIKE %s ESCAPE \'{esc}\'')
        params.append(pat)

    qt = (q_global or "").strip()
    if qt:
        pat = "%" + _like_escape(qt, esc) + "%"
        ors: List[str] = []
        for col in sorted(allowed):
            if not IDENT_RE.match(col):
                continue
            ors.append(f'"{col}"::text ILIKE %s ESCAPE \'{esc}\'')
            params.append(pat)
        if ors:
            parts.append("(" + " OR ".join(ors) + ")")

    if not parts:
        return "", []
    return " AND ".join(parts), params


def _order_expr_for_column(meta_row: Dict[str, Any]) -> str:
    col = meta_row["name"]
    if meta_row.get("udt_name") == "vector":
        return f'"{col}"::text'
    return f'"{col}"'


def build_order_clause(
    meta: List[Dict[str, Any]],
    sort_by: Optional[str],
    sort_order_raw: Optional[str],
) -> str:
    by_name = {r["name"]: r for r in meta}
    names_order = [r["name"] for r in meta]
    order = (sort_order_raw or "asc").lower()
    if order not in ("asc", "desc"):
        order = "asc"
    col = sort_by
    if (
        not col
        or col not in by_name
        or not IDENT_RE.match(col)
    ):
        if "id" in by_name:
            col = "id"
            order = "desc"
        elif names_order:
            col = names_order[0]
            order = "asc"
        else:
            return "ORDER BY 1 ASC NULLS LAST"
    expr = _order_expr_for_column(by_name[col])
    return f"ORDER BY {expr} {order.upper()} NULLS LAST"


def normalize_sort_order_param(raw: Optional[str]) -> Optional[str]:
    if not raw or not str(raw).strip():
        return None
    x = str(raw).strip().lower()
    return x if x in ("asc", "desc") else None


def fetch_table_page(
    cur: Any,
    schema: str,
    table_name: str,
    meta: List[Dict[str, Any]],
    filters_raw: Optional[str],
    q: Optional[str],
    sort_by: Optional[str],
    sort_order: Optional[str],
    limit: int,
    offset: int,
    *,
    with_total: bool,
) -> tuple[List[Dict[str, Any]], Optional[int]]:
    fd = (
        parse_filters_query(filters_raw)
        if filters_raw and str(filters_raw).strip()
        else {}
    )
    where_sql, wparams = build_fuzzy_where_clause(meta, fd, q)
    order_clause = build_order_clause(meta, sort_by, sort_order)
    base_from = f'FROM "{schema}"."{table_name}"'
    frag = (" WHERE " + where_sql) if where_sql else ""

    total: Optional[int] = None
    if with_total:
        cur.execute(f"SELECT COUNT(*) AS _cnt {base_from}{frag}", tuple(wparams))
        total = int(cur.fetchone()["_cnt"])

    cur.execute(
        f"SELECT * {base_from}{frag} {order_clause} LIMIT %s OFFSET %s",
        (*wparams, limit, offset),
    )
    rows = [dict(r) for r in cur.fetchall()]
    return rows, total


# ── HTTP：表与导航 ───────────────────────────────────────────────────────────


@app.get("/")
def root():
    return {"status": "ok", "project": "AI·ERP"}


@app.get("/health")
def health():
    return {"status": "ok"}


@app.get("/tables")
def list_tables():
    with db_conn() as conn:
        with conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor) as cur:
            cur.execute(
                """
                SELECT table_schema AS schema, table_name AS table
                FROM information_schema.tables
                WHERE table_schema IN ('business', 'ai_assistant')
                  AND table_type = 'BASE TABLE'
                ORDER BY table_schema, table_name
                """
            )
            rows = cur.fetchall()
        return [dict(r) for r in rows]


def _schema_by_table_name(conn) -> dict:
    with conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor) as cur:
        cur.execute(
            """
            SELECT table_name, table_schema
            FROM information_schema.tables
            WHERE table_schema IN ('meta', 'business', 'ai_assistant') AND table_type = 'BASE TABLE'
            """
        )
        m = {}
        for r in cur.fetchall():
            m[r["table_name"]] = r["table_schema"]
        return m


@app.get("/nav/tables")
def list_nav_tables():
    with db_conn() as conn:
        schema_map = _schema_by_table_name(conn)
        with conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor) as cur:
            cur.execute(
                """
                SELECT table_name, table_cn, nav_section, nav_order, nav_parent_table
                FROM meta.meta_tables
                WHERE nav_section IS NOT NULL
                  AND nav_parent_table IS NULL
                ORDER BY nav_section, nav_order, table_name
                """
            )
            rows = cur.fetchall()
        out = {"design": [], "prep": [], "run": []}
        for r in rows:
            sec = r.get("nav_section")
            if sec not in out:
                continue
            sch = schema_map.get(r["table_name"])
            if not sch or sch == "meta":
                continue
            out[sec].append(
                {
                    "schema": sch,
                    "table_name": r["table_name"],
                    "table_cn": r["table_cn"],
                    "nav_order": r["nav_order"],
                    "nav_parent_table": r["nav_parent_table"],
                }
            )
        return out


@app.get("/api/user-menu")
def api_user_menu(
    user_id: int = Query(..., description="employees.id，登录用户 id"),
):
    """
    权限驱动的导航：与 /nav/tables 结构相同，按 can_read 过滤；
    admin/super_admin 可见全部 business 导航项。
    仅包含 nav_parent_table 为空的项（侧栏只展示主表，子表在详情内查看）。
    """
    try:
        with db_conn() as conn:
            with conn.cursor() as cur:
                roles = _employee_role_codes(cur, user_id)
            privileged = any(r in ("admin", "super_admin") for r in roles)
            schema_map = _schema_by_table_name(conn)
            with conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor) as cur:
                cur.execute(
                    """
                    SELECT table_name, table_cn, nav_section, nav_order, nav_parent_table
                    FROM meta.meta_tables
                    WHERE nav_section IS NOT NULL
                      AND nav_parent_table IS NULL
                    ORDER BY nav_section, nav_order, table_name
                    """
                )
                rows = cur.fetchall()
            out: Dict[str, List[Dict[str, Any]]] = {"design": [], "prep": [], "run": []}
            for r in rows:
                sec = r.get("nav_section")
                if sec not in out:
                    continue
                sch = schema_map.get(r["table_name"])
                if not sch or sch == "meta":
                    continue
                if sch != "business":
                    continue
                full_tn = f"business.{r['table_name']}"
                if not privileged and not user_can_read_table(conn, user_id, full_tn):
                    continue
                out[sec].append(
                    {
                        "schema": sch,
                        "table_name": r["table_name"],
                        "table_cn": r["table_cn"],
                        "nav_order": r["nav_order"],
                        "nav_parent_table": r["nav_parent_table"],
                    }
                )
        return {
            "design": out["design"],
            "prep": out["prep"],
            "run": out["run"],
            "role": roles,
        }
    except Exception as e:
        logger.exception("api/user-menu: %s", e)
        raise HTTPException(status_code=500, detail="菜单加载失败") from e


def _apply_dashboard_user_id(sql: str, user_id: int) -> str:
    """将卡片 SQL 中的 :user_id 替换为当前登录员工 id（整数）。"""
    return (sql or "").replace(":user_id", str(int(user_id)))


@app.get("/api/dashboard")
def get_dashboard(user_id: int = Query(..., description="employees.id")):
    """
    个人卡片优先；无则角色默认卡片。
    每张卡片的 query_sql 经 execute_sql 执行（仅 SELECT，R1–R7）。
    """
    results: List[Dict[str, Any]] = []
    try:
        with db_conn() as conn:
            with conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor) as cur:
                cur.execute(
                    """
                    SELECT id, employee_id, role_code, title, card_type, query_sql, description,
                           display_order, card_config, is_active, created_by, created_at, updated_at
                    FROM business.dashboard_cards
                    WHERE employee_id = %s AND is_active = true AND deleted_at IS NULL
                    ORDER BY display_order, id
                    """,
                    (user_id,),
                )
                cards = cur.fetchall()

            if not cards:
                with conn.cursor() as cur:
                    cur.execute(
                        """
                        SELECT DISTINCT rg.role_code
                        FROM business.user_role_grants rg
                        JOIN business.user_accounts ua ON ua.id = rg.user_account_id
                        WHERE ua.employee_id = %s AND ua.is_active = true
                        """,
                        (user_id,),
                    )
                    role_codes = [r[0] for r in cur.fetchall() if r and r[0]]
                if role_codes:
                    with conn.cursor(
                        cursor_factory=psycopg2.extras.RealDictCursor
                    ) as cur:
                        cur.execute(
                            """
                            SELECT id, employee_id, role_code, title, card_type, query_sql, description,
                                   display_order, card_config, is_active, created_by, created_at, updated_at
                            FROM business.dashboard_cards
                            WHERE role_code = ANY(%s) AND employee_id IS NULL
                              AND is_active = true AND deleted_at IS NULL
                            ORDER BY display_order, id
                            """,
                            (list(role_codes),),
                        )
                        cards = cur.fetchall()
                else:
                    cards = []

            for card in cards:
                raw_sql = card["query_sql"] or ""
                sql = _apply_dashboard_user_id(raw_sql, user_id)
                cfg = card.get("card_config")
                if isinstance(cfg, str):
                    try:
                        cfg = json.loads(cfg)
                    except Exception:
                        cfg = {}
                elif cfg is None or not isinstance(cfg, dict):
                    cfg = {}
                base: Dict[str, Any] = {
                    "id": card["id"],
                    "title": card["title"],
                    "card_type": card["card_type"],
                    "description": card.get("description"),
                    "card_config": cfg,
                }
                out = execute_sql(sql, conn=conn, user_id=user_id)
                if out.get("success"):
                    base["data"] = out.get("rows") or []
                    base["error"] = None
                else:
                    base["data"] = []
                    base["error"] = out.get("error") or "执行失败"
                results.append(base)

        return jsonable_encoder({"cards": results})
    except Exception as e:
        logger.exception("api/dashboard: %s", e)
        raise HTTPException(status_code=500, detail="您还未创建个人仪表盘。") from e


@app.get("/meta/table-cn/{table_name}")
def get_table_cn(table_name: str):
    """从 meta.meta_tables 取中文表名（用于未出现在导航列表中的子表等）。"""
    assert_ident(table_name, "table_name")
    with db_conn() as conn:
        with conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor) as cur:
            cur.execute(
                """
                SELECT table_name, table_cn
                FROM meta.meta_tables
                WHERE table_name = %s
                """,
                (table_name,),
            )
            r = cur.fetchone()
    if not r:
        return {"table_name": table_name, "table_cn": None}
    cn = r.get("table_cn")
    return {"table_name": r["table_name"], "table_cn": cn if cn else None}


@app.get("/table-columns/{table_name}")
def get_table_column_labels(table_name: str):
    assert_ident(table_name, "table_name")
    with db_conn() as conn:
        schema_map = _schema_by_table_name(conn)
        sch = schema_map.get(table_name)
        if sch == "meta":
            raise HTTPException(
                status_code=403,
                detail="系统表不可通过界面访问",
            )
        with conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor) as cur:
            cur.execute(
                """
                SELECT field_name, COALESCE(menu_label_cn, field_name) AS menu_label_cn,
                       COALESCE(menu_label_en, field_name) AS menu_label_en,
                       data_type
                FROM meta.meta_fields
                WHERE table_name = %s
                ORDER BY id
                """,
                (table_name,),
            )
            cols = cur.fetchall()
        return {"columns": [dict(c) for c in cols]}


@app.get("/table/{schema}/{table_name}")
def get_table(
    schema: str,
    table_name: str,
    limit: int = 500,
    offset: int = Query(0, ge=0, le=1_000_000),
    sort_by: Optional[str] = Query(None),
    sort_order: Optional[str] = Query(None),
    filters: Optional[str] = Query(None),
    q: Optional[str] = Query(None),
):
    if schema not in ALLOWED_SCHEMAS:
        raise HTTPException(status_code=400, detail="Invalid schema")
    _forbid_meta_table_api(schema)
    assert_ident(table_name, "table_name")
    lim = min(max(1, limit), 5000)
    off = max(0, min(offset, 1_000_000))
    so = normalize_sort_order_param(sort_order)
    with db_conn() as conn:
        with conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor) as cur:
            try:
                meta = fetch_column_meta(cur, schema, table_name)
                if not meta:
                    raise HTTPException(status_code=404, detail="Table not found")
                rows, total = fetch_table_page(
                    cur,
                    schema,
                    table_name,
                    meta,
                    filters,
                    q,
                    sort_by,
                    so,
                    lim,
                    off,
                    with_total=True,
                )
            except HTTPException:
                raise
            except Exception as e:
                raise HTTPException(status_code=400, detail=str(e)) from e
            return {"rows": rows, "total": total}


def _export_cell_value(val: Any, col_meta: Optional[Dict[str, Any]]) -> Any:
    if val is None:
        return None
    udt = (col_meta or {}).get("udt_name")
    if udt == "vector":
        return "(向量，已省略)"
    if isinstance(val, uuid.UUID):
        return str(val)
    if isinstance(val, (bytes, memoryview)):
        return "(二进制)"
    if isinstance(val, (dict, list)):
        return json.dumps(val, ensure_ascii=False)
    if isinstance(val, Decimal):
        return str(val)
    if isinstance(val, datetime):
        if val.tzinfo is not None:
            return val.astimezone(timezone.utc).replace(tzinfo=None)
        return val
    if isinstance(val, date):
        return val
    if isinstance(val, time):
        if val.tzinfo is not None:
            return val.replace(tzinfo=None)
        return val
    if isinstance(val, bool):
        return val
    if isinstance(val, (int, float, str)):
        return val
    return str(val)


def _resolve_export_columns(
    meta: List[Dict[str, Any]], cols_param: Optional[str]
) -> List[str]:
    all_names = [r["name"] for r in meta]
    name_set = set(all_names)
    if not cols_param or not cols_param.strip():
        return all_names
    out: List[str] = []
    seen = set()
    for raw in cols_param.split(","):
        c = raw.strip()
        if not c or c in seen:
            continue
        if not IDENT_RE.match(c) or c not in name_set:
            continue
        seen.add(c)
        out.append(c)
    return out if out else all_names


@app.get("/export/xlsx/{schema}/{table_name}")
def export_table_xlsx(
    schema: str,
    table_name: str,
    limit: int = 5000,
    cols: Optional[str] = None,
    filters: Optional[str] = Query(None),
    q: Optional[str] = Query(None),
    sort_by: Optional[str] = Query(None),
    sort_order: Optional[str] = Query(None),
):
    try:
        from openpyxl import Workbook
    except ImportError as e:
        raise HTTPException(
            status_code=503,
            detail="缺少 openpyxl，无法导出。请 pip install openpyxl",
        ) from e

    if schema not in ALLOWED_SCHEMAS:
        raise HTTPException(status_code=400, detail="Invalid schema")
    _forbid_meta_table_api(schema)
    assert_ident(table_name, "table_name")
    lim = min(max(1, limit), 10000)
    so = normalize_sort_order_param(sort_order)

    with db_conn() as conn:
        with conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor) as cur:
            meta = fetch_column_meta(cur, schema, table_name)
            if not meta:
                raise HTTPException(status_code=404, detail="Table not found")

            col_order = _resolve_export_columns(meta, cols)
            name_to_meta = {r["name"]: r for r in meta}

            cur.execute(
                """
                SELECT field_name, COALESCE(menu_label_cn, field_name) AS menu_label_cn
                FROM meta.meta_fields
                WHERE table_name = %s
                """,
                (table_name,),
            )
            label_map = {r["field_name"]: r["menu_label_cn"] for r in cur.fetchall()}
            headers = [label_map.get(c, c) for c in col_order]

            try:
                db_rows_dicts, _ = fetch_table_page(
                    cur,
                    schema,
                    table_name,
                    meta,
                    filters,
                    q,
                    sort_by,
                    so,
                    lim,
                    0,
                    with_total=False,
                )
            except HTTPException:
                raise
            except Exception as e:
                raise HTTPException(status_code=400, detail=str(e)) from e

    wb = Workbook()
    ws = wb.active
    ws.title = "data"[:31]
    ws.append(headers)
    for d in db_rows_dicts:
        ws.append(
            [_export_cell_value(d.get(c), name_to_meta.get(c)) for c in col_order]
        )

    buf = io.BytesIO()
    wb.save(buf)
    body = buf.getvalue()
    fname = f"{table_name}.xlsx"
    return Response(
        content=body,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{fname}"'},
    )


@app.post("/table/{schema}/{table_name}")
def table_insert(schema: str, table_name: str, body: Dict[str, Any] = Body(...)):
    if schema not in ALLOWED_SCHEMAS:
        raise HTTPException(status_code=400, detail="Invalid schema")
    _forbid_meta_table_api(schema)
    assert_ident(schema, "schema")
    assert_ident(table_name, "table_name")
    with db_conn() as conn:
        with conn:
            with conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor) as cur:
                meta = fetch_column_meta(cur, schema, table_name)
                if not meta:
                    raise HTTPException(status_code=404, detail="Table not found")
                allowed = {r["name"]: r for r in writable_columns(meta)}
                cols = []
                vals = []
                for k, v in body.items():
                    if k not in allowed:
                        continue
                    r = allowed[k]
                    if k == "id" and (v is None or v == ""):
                        continue
                    cols.append(k)
                    vals.append(
                        coerce_cell(v, r["udt_name"], r["data_type"], r["is_nullable"])
                    )
                if not cols:
                    raise HTTPException(status_code=400, detail="No writable columns in body")
                qcols = ",".join(f'"{c}"' for c in cols)
                ph = ",".join(["%s"] * len(cols))
                sql = f'INSERT INTO "{schema}"."{table_name}" ({qcols}) VALUES ({ph}) RETURNING *'
                try:
                    cur.execute(sql, vals)
                    row = cur.fetchone()
                except Exception as e:
                    raise HTTPException(status_code=400, detail=str(e)) from e
        return {"row": jsonable_encoder(dict(row))}


@app.patch("/table/{schema}/{table_name}/{row_id}")
def table_update(schema: str, table_name: str, row_id: int, body: Dict[str, Any] = Body(...)):
    if schema not in ALLOWED_SCHEMAS:
        raise HTTPException(status_code=400, detail="Invalid schema")
    _forbid_meta_table_api(schema)
    assert_ident(schema, "schema")
    assert_ident(table_name, "table_name")
    with db_conn() as conn:
        with conn:
            with conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor) as cur:
                meta = fetch_column_meta(cur, schema, table_name)
                if not meta or not table_has_id(meta):
                    raise HTTPException(status_code=400, detail="Table must have integer id column")
                allowed = {r["name"]: r for r in writable_columns(meta)}
                sets = []
                vals = []
                for k, v in body.items():
                    if k == "id":
                        continue
                    if k not in allowed:
                        continue
                    r = allowed[k]
                    sets.append(f'"{k}" = %s')
                    vals.append(
                        coerce_cell(v, r["udt_name"], r["data_type"], r["is_nullable"])
                    )
                if not sets:
                    raise HTTPException(status_code=400, detail="No updatable fields")
                vals.append(row_id)
                sql = f'UPDATE "{schema}"."{table_name}" SET {", ".join(sets)} WHERE id = %s RETURNING *'
                try:
                    cur.execute(sql, vals)
                    row = cur.fetchone()
                    if not row:
                        raise HTTPException(status_code=404, detail="Row not found")
                except HTTPException:
                    raise
                except Exception as e:
                    raise HTTPException(status_code=400, detail=str(e)) from e
        return {"row": jsonable_encoder(dict(row))}


@app.delete("/table/{schema}/{table_name}/{row_id}")
def table_delete(schema: str, table_name: str, row_id: int):
    if schema not in ALLOWED_SCHEMAS:
        raise HTTPException(status_code=400, detail="Invalid schema")
    _forbid_meta_table_api(schema)
    assert_ident(schema, "schema")
    assert_ident(table_name, "table_name")
    with db_conn() as conn:
        with conn:
            with conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor) as cur:
                meta = fetch_column_meta(cur, schema, table_name)
                if not meta or not table_has_id(meta):
                    raise HTTPException(status_code=400, detail="Table must have integer id column")
                try:
                    cur.execute(
                        f'DELETE FROM "{schema}"."{table_name}" WHERE id = %s RETURNING *',
                        (row_id,),
                    )
                    row = cur.fetchone()
                    if not row:
                        raise HTTPException(status_code=404, detail="Row not found")
                except HTTPException:
                    raise
                except Exception as e:
                    raise HTTPException(status_code=400, detail=str(e)) from e
        return {"row": jsonable_encoder(dict(row))}


# ── 认证（简单登录，无 JWT）──────────────────────────────────────────────────


class LoginRequest(BaseModel):
    login_name: str
    password: str


def _row_to_user_dict(row: tuple) -> Dict[str, Any]:
    """employees 查询行 → 前端 user 对象（id 为 employee id）。"""
    eid, emp_code, name, department, role_tags = row
    return {
        "id": int(eid),
        "emp_code": emp_code or "",
        "name": name or "",
        "department": department or "",
        "role_tags": role_tags if role_tags is not None else "",
    }


def _fetch_employee_user(cur, employee_id: int) -> Optional[Dict[str, Any]]:
    cur.execute(
        """
        SELECT id, emp_code, name, department, role_tags
        FROM business.employees
        WHERE id = %s
        """,
        (employee_id,),
    )
    row = cur.fetchone()
    if not row:
        return None
    return _row_to_user_dict(row)


@app.post("/auth/login")
async def auth_login(req: LoginRequest):
    """
    简单登录：验证 login_name + password，返回 user（employee 维度）。
    不做 JWT；前端保存 user_id 后随请求携带。
    """
    name = (req.login_name or "").strip()
    if not name or not req.password:
        return {"success": False, "error": "用户名或密码错误"}

    try:
        with db_conn() as conn:
            with conn.cursor() as cur:
                cur.execute(
                    """
                    SELECT id, employee_id, password_hash, is_active
                    FROM business.user_accounts
                    WHERE LOWER(login_name) = LOWER(%s)
                    """,
                    (name,),
                )
                row = cur.fetchone()
                if not row or not row[2]:
                    return {"success": False, "error": "用户名或密码错误"}
                if not row[3]:
                    return {"success": False, "error": "用户名或密码错误"}

                ph = row[2]
                if isinstance(ph, memoryview):
                    ph = ph.tobytes().decode("utf-8", errors="replace")
                else:
                    ph = str(ph)

                try:
                    ok = bcrypt.checkpw(
                        req.password.encode("utf-8"),
                        ph.encode("utf-8"),
                    )
                except ValueError:
                    ok = False
                if not ok:
                    return {"success": False, "error": "用户名或密码错误"}

                ua_id = int(row[0])
                employee_id = int(row[1])

                user = _fetch_employee_user(cur, employee_id)
                if not user:
                    return {"success": False, "error": "用户名或密码错误"}

                cur.execute(
                    """
                    UPDATE business.user_accounts
                    SET last_login_at = NOW(), updated_at = NOW()
                    WHERE id = %s
                    """,
                    (ua_id,),
                )
                try:
                    cur.execute(
                        """
                        UPDATE business.employees
                        SET last_login_at = NOW()
                        WHERE id = %s
                        """,
                        (employee_id,),
                    )
                except Exception:
                    pass

            conn.commit()
    except Exception as e:
        logger.exception("auth_login: %s", e)
        return {"success": False, "error": "登录失败，请稍后重试"}

    return {"success": True, "user": user}


@app.get("/auth/me")
async def auth_me(user_id: int = Query(..., description="employees.id")):
    """根据 user_id（员工 id）恢复登录态。"""
    try:
        with db_conn() as conn:
            with conn.cursor() as cur:
                user = _fetch_employee_user(cur, user_id)
    except Exception as e:
        logger.exception("auth_me: %s", e)
        raise HTTPException(status_code=500, detail="查询失败") from e

    if not user:
        raise HTTPException(status_code=404, detail="用户不存在")
    return {"user": user}


def _employee_role_codes(cur, employee_id: int) -> List[str]:
    cur.execute(
        """
        SELECT DISTINCT rg.role_code
        FROM business.user_role_grants rg
        JOIN business.user_accounts ua ON ua.id = rg.user_account_id
        WHERE ua.employee_id = %s AND ua.is_active = true
        """,
        (employee_id,),
    )
    return [str(r[0]) for r in cur.fetchall() if r and r[0]]


@app.get("/approvals/pending")
async def approvals_pending(
    user_id: int = Query(..., description="employees.id，用于筛选与我相关的待办"),
):
    """待审批列表：本人发起或指派给本人的 pending；admin/super_admin 可见全部 pending。"""
    try:
        with db_conn() as conn:
            with conn.cursor() as cur:
                roles = _employee_role_codes(cur, user_id)
            privileged = any(r in ("admin", "super_admin") for r in roles)
            with conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor) as cur:
                if privileged:
                    cur.execute(
                        """
                        SELECT a.id, a.approval_no, a.target_type, a.target_id, a.target_no,
                               a.requested_by, a.requested_at, a.approver_id, a.status,
                               e1.name AS requested_by_name,
                               e2.name AS approver_name
                        FROM business.approvals a
                        LEFT JOIN business.employees e1 ON e1.id = a.requested_by
                        LEFT JOIN business.employees e2 ON e2.id = a.approver_id
                        WHERE a.deleted_at IS NULL AND a.status = 'pending'
                        ORDER BY a.requested_at DESC
                        """
                    )
                else:
                    cur.execute(
                        """
                        SELECT a.id, a.approval_no, a.target_type, a.target_id, a.target_no,
                               a.requested_by, a.requested_at, a.approver_id, a.status,
                               e1.name AS requested_by_name,
                               e2.name AS approver_name
                        FROM business.approvals a
                        LEFT JOIN business.employees e1 ON e1.id = a.requested_by
                        LEFT JOIN business.employees e2 ON e2.id = a.approver_id
                        WHERE a.deleted_at IS NULL AND a.status = 'pending'
                          AND (a.requested_by = %s OR a.approver_id = %s)
                        ORDER BY a.requested_at DESC
                        """,
                        (user_id, user_id),
                    )
                rows = cur.fetchall() or []
    except Exception as e:
        logger.exception("approvals/pending: %s", e)
        raise HTTPException(status_code=500, detail="查询失败") from e
    return {"items": [dict(r) for r in rows]}


class ApprovalDecideRequest(BaseModel):
    user_id: int = Field(..., description="employees.id，审批人")
    decision: str = Field(..., description="approved 或 rejected")
    comment: Optional[str] = None


@app.post("/approvals/{approval_id}/decide")
async def approval_decide(approval_id: int, body: ApprovalDecideRequest):
    """审批通过/驳回，并回写合同或采购订单状态。"""
    d = (body.decision or "").strip().lower()
    if d not in ("approved", "rejected"):
        raise HTTPException(status_code=400, detail="decision 须为 approved 或 rejected")
    comment = (body.comment or "").strip() or None
    try:
        with db_conn() as conn:
            with conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor) as cur:
                cur.execute(
                    """
                    SELECT id, target_type, target_id, status, requested_by, approver_id
                    FROM business.approvals
                    WHERE id = %s AND deleted_at IS NULL
                    """,
                    (approval_id,),
                )
                row = cur.fetchone()
                if not row:
                    raise HTTPException(status_code=404, detail="审批单不存在")
                if (row.get("status") or "") != "pending":
                    raise HTTPException(status_code=400, detail="该审批单已处理")

                roles = _employee_role_codes(cur, body.user_id)
                is_admin = any(r in ("admin", "super_admin") for r in roles)
                aid = row.get("approver_id")
                if not is_admin:
                    if aid is not None and int(aid) != int(body.user_id):
                        raise HTTPException(status_code=403, detail="无权审批该单据")
                    if aid is None:
                        raise HTTPException(
                            status_code=403,
                            detail="未指定审批人时仅管理员可审批",
                        )

                new_ap_status = "approved" if d == "approved" else "rejected"
                cur.execute(
                    """
                    UPDATE business.approvals
                    SET status = %s,
                        decision_at = NOW(),
                        comment = %s,
                        approver_id = COALESCE(approver_id, %s),
                        updated_at = NOW()
                    WHERE id = %s AND status = 'pending'
                    RETURNING id
                    """,
                    (new_ap_status, comment, body.user_id, approval_id),
                )
                if not cur.fetchone():
                    raise HTTPException(status_code=400, detail="更新审批单失败")

                tt = str(row.get("target_type") or "")
                tid = row.get("target_id")
                if tid is None:
                    conn.rollback()
                    raise HTTPException(status_code=500, detail="审批单缺少 target_id")
                tid = int(tid)

                if d == "approved":
                    if tt == "contract":
                        cur.execute(
                            "UPDATE business.contracts SET status = 'active' WHERE id = %s",
                            (tid,),
                        )
                    elif tt == "purchase_order":
                        # 与合同审批通过后均为 active，下游只认一种「已生效」状态
                        cur.execute(
                            "UPDATE business.purchase_orders SET status = 'active' WHERE id = %s",
                            (tid,),
                        )
                else:
                    if tt == "contract":
                        cur.execute(
                            "UPDATE business.contracts SET status = 'rejected' WHERE id = %s",
                            (tid,),
                        )
                    elif tt == "purchase_order":
                        cur.execute(
                            "UPDATE business.purchase_orders SET status = 'rejected' WHERE id = %s",
                            (tid,),
                        )
            conn.commit()
    except HTTPException:
        raise
    except Exception as e:
        logger.exception("approvals/decide: %s", e)
        raise HTTPException(status_code=500, detail="处理失败") from e
    return {"ok": True, "approval_id": approval_id, "status": new_ap_status}


# ── AI Chat（直连 LLM，与前端 AiPanel 协议对齐）──────────────────────────────


class ChatMessage(BaseModel):
    role: str
    content: str


class WriteConfirmRequest(BaseModel):
    """确认执行 write：传入 preview 返回的 payload；兼容旧字段 frame_id（已停用）。"""

    model_config = ConfigDict(populate_by_name=True)

    payload: Optional[Dict[str, Any]] = None
    frame_id: Optional[int] = None
    confirmed: Optional[bool] = None
    session_id: Optional[int] = None
    user_id: Optional[int] = Field(
        default=None,
        description=(
            "employees.id，用于 R7 与审批；应与发起预览时为同一登录用户，"
            "否则可能出现预览通过但确认阶段因角色不同被拒"
        ),
    )


class ChatLogMessageRequest(BaseModel):
    """将一条对话记入 ai_chat（用于前端取消/异常等仅客户端发生的收尾文案）。"""

    session_id: int
    content: str
    role: str = "assistant"


class ChatRequest(BaseModel):
    model_config = ConfigDict(populate_by_name=True)

    messages: List[ChatMessage]
    images: Optional[List[str]] = None
    files: Optional[List[str]] = None
    confirm: Optional[bool] = False
    pending_token: Optional[str] = Field(default=None, alias="pendingToken")
    stream: Optional[bool] = False
    invite_expert: Optional[bool] = Field(default=False, alias="inviteExpert")
    session_id: Optional[int] = None
    user_id: Optional[int] = None


def _validate_write_payload_to_batch(pl: Dict[str, Any]) -> List[Dict[str, Any]]:
    """将前端 payload 规范为 confirm_batch_sql 使用的 batch。"""
    raw = pl.get("batch")
    if isinstance(raw, list) and len(raw) > 0:
        out: List[Dict[str, Any]] = []
        for i, item in enumerate(raw):
            if not isinstance(item, dict):
                raise HTTPException(status_code=400, detail=f"batch[{i}] 须为对象")
            if not (item.get("sql") or "").strip():
                raise HTTPException(status_code=400, detail=f"batch[{i}] 缺少 sql")
            if not isinstance(item.get("params", []), list):
                raise HTTPException(status_code=400, detail=f"batch[{i}].params 必须为数组")
            out.append(
                {
                    "sql": (item.get("sql") or "").strip(),
                    "params": list(item.get("params") or []),
                }
            )
        return out
    sql = pl.get("sql")
    params = pl.get("params", [])
    if not sql or not isinstance(sql, str):
        raise HTTPException(status_code=400, detail="缺少 sql 或 batch")
    if not isinstance(params, list):
        raise HTTPException(status_code=400, detail="params 必须为数组")
    if not sql.strip():
        raise HTTPException(status_code=400, detail="缺少 sql")
    return [{"sql": sql.strip(), "params": params}]


def _merge_tool_context_into_payload(
    payload: Any,
    tool_block: Dict[str, str],
) -> Dict[str, Any]:
    """预览 payload 附带原始 [TOOL] 意图，供 /ai/confirm 失败时专家恢复引用（避免只修主表丢掉子表）。"""
    if isinstance(payload, dict):
        base = dict(payload)
    else:
        base = {}
    base["tool_context"] = dict(tool_block)
    return base


def verify_batch_matches_tool_intent(
    batch: List[Dict[str, Any]],
    tool_context: Optional[Dict[str, Any]],
) -> Dict[str, Any]:
    """
    执行前检查：若意图明确包含合同行项目，则 batch 须出现 contract_items 相关写入。
    避免主表 INSERT 已成功、子表从未生成却显示「完成任务」。
    """
    if not tool_context or not isinstance(tool_context, dict):
        return {"ok": True, "reason": ""}
    details = str(tool_context.get("details") or "")
    action = str(tool_context.get("action") or "").lower()
    intent = str(tool_context.get("intent") or "")
    blob = f"{action} {intent} {details}"
    if "合同" not in blob and "contract" not in blob.lower():
        return {"ok": True, "reason": ""}
    line_hints = ("数量", "米", "产品", "单价", "金额", "明细", "行项目")
    if not any(h in details for h in line_hints):
        return {"ok": True, "reason": ""}
    sql_joined = " ".join((x.get("sql") or "").lower() for x in batch)
    if "contract_items" in sql_joined:
        return {"ok": True, "reason": ""}
    return {
        "ok": False,
        "reason": (
            "意图（details）包含合同行项目信息，但当前待执行 SQL 未写入 business.contract_items；"
            "请生成含主表与子表的完整 batch（可用 CTE 多语句或顺序 INSERT）。"
        ),
    }


def _batch_from_write_preview(expert_out: Dict[str, Any]) -> Optional[List[Dict[str, Any]]]:
    """从专家返回的预览结果中提取可执行的 batch。"""
    if not _is_write_preview_result(expert_out):
        return None
    pl = expert_out.get("payload") if isinstance(expert_out.get("payload"), dict) else {}
    raw = pl.get("batch")
    if not isinstance(raw, list) or len(raw) == 0:
        return None
    out: List[Dict[str, Any]] = []
    for item in raw:
        if not isinstance(item, dict):
            return None
        s = (item.get("sql") or "").strip()
        if not s:
            return None
        out.append({"sql": s, "params": list(item.get("params") or [])})
    return out


async def _confirm_writes_with_verify_and_recovery(
    batch: List[Dict[str, Any]],
    tool_context: Optional[Dict[str, Any]] = None,
    user_id: Optional[int] = None,
    session_id: Optional[int] = None,
) -> Dict[str, Any]:
    """执行前意图检查 → 执行 → 校验；失败则调用专家重写 SQL，最多共 MAX_WRITE_CONFIRM_ATTEMPTS 次执行尝试。"""
    current: List[Dict[str, Any]] = [dict(x) for x in batch]
    recovery_rounds = 0

    for attempt in range(MAX_WRITE_CONFIRM_ATTEMPTS):
        intent_chk = verify_batch_matches_tool_intent(current, tool_context)
        if not intent_chk.get("ok"):
            failure_reason = intent_chk.get("reason") or "与原始意图不一致"
            if attempt >= MAX_WRITE_CONFIRM_ATTEMPTS - 1:
                return {
                    "success": False,
                    "error": (
                        f"任务失败：已尝试 {MAX_WRITE_CONFIRM_ATTEMPTS} 次仍未满足意图校验。"
                        f"最后原因：{failure_reason}"
                    ),
                    "task_failed": True,
                    "recovery_rounds": recovery_rounds,
                    "last_error": failure_reason,
                }
            expert_out = await _call_tool_expert_recover_after_failed_confirm(
                failure_reason,
                current,
                None,
                tool_context,
                user_id=user_id,
                session_id=session_id,
            )
            new_batch = _batch_from_write_preview(expert_out)
            if not new_batch:
                err = expert_out.get("error") or "专家未返回可执行的写入预览"
                log_incident(
                    stage="confirm",
                    error_message=str(err),
                    session_id=session_id,
                    user_id=user_id,
                    generated_sql=json.dumps(current, ensure_ascii=False, default=str)[:5000],
                    context={
                        "phase": "intent_verify",
                        "failure_reason": (failure_reason or "")[:800],
                    },
                )
                return {
                    "success": False,
                    "error": f"任务失败：{err}",
                    "task_failed": True,
                    "recovery_rounds": recovery_rounds,
                    "last_error": err,
                }
            current = new_batch
            recovery_rounds += 1
            continue

        def _run_one() -> Dict[str, Any]:
            with db_conn() as conn:
                try:
                    return confirm_batch_sql(current, conn, user_id=user_id)
                except Exception as e:
                    logger.exception("ai/confirm confirm_batch_sql: %s", e)
                    log_incident(
                        stage="execute",
                        error_message=str(e),
                        session_id=session_id,
                        user_id=user_id,
                        generated_sql=json.dumps(current, ensure_ascii=False, default=str)[:5000],
                        context={"fn": "confirm_batch_sql", "phase": "confirm_path"},
                    )
                    return {"success": False, "error": f"执行异常：{e}"}

        confirm_result = await asyncio.to_thread(_run_one)

        if confirm_result.get("success"):
            v = verify_confirmed_writes(current, confirm_result)
            if v.get("ok"):
                return {
                    "success": True,
                    "results": confirm_result.get("results"),
                    "count": confirm_result.get("count"),
                    "verify_ok": True,
                    "recovery_rounds": recovery_rounds,
                }
            failure_reason = v.get("reason") or "执行后校验未通过"
        else:
            failure_reason = confirm_result.get("error") or "执行失败"

        if attempt >= MAX_WRITE_CONFIRM_ATTEMPTS - 1:
            return {
                "success": False,
                "error": (
                    f"任务失败：已尝试 {MAX_WRITE_CONFIRM_ATTEMPTS} 次执行仍未成功。"
                    f"最后原因：{failure_reason}"
                ),
                "task_failed": True,
                "recovery_rounds": recovery_rounds,
                "last_error": failure_reason,
            }

        expert_out = await _call_tool_expert_recover_after_failed_confirm(
            failure_reason,
            current,
            confirm_result,
            tool_context,
            user_id=user_id,
            session_id=session_id,
        )
        new_batch = _batch_from_write_preview(expert_out)
        if not new_batch:
            err = expert_out.get("error") or "专家未返回可执行的写入预览"
            log_incident(
                stage="confirm",
                error_message=str(err),
                session_id=session_id,
                user_id=user_id,
                generated_sql=json.dumps(current, ensure_ascii=False, default=str)[:5000],
                context={
                    "phase": "post_exec",
                    "failure_reason": (failure_reason or "")[:800],
                },
            )
            return {
                "success": False,
                "error": f"任务失败：{err}",
                "task_failed": True,
                "recovery_rounds": recovery_rounds,
                "last_error": err,
            }
        current = new_batch
        recovery_rounds += 1

    return {
        "success": False,
        "error": "任务失败：未预期结束",
        "task_failed": True,
        "recovery_rounds": recovery_rounds,
    }


def _sse_line(obj: Dict[str, Any]) -> bytes:
    return f"data: {json.dumps(obj, ensure_ascii=False)}\n\n".encode("utf-8")


def _assistant_line_from_write_confirm_result(r: Dict[str, Any]) -> str:
    """兼容旧逻辑：失败时的简短展示（主路径已改为 TOOL_RESULT + 主助手续写）。"""
    if r.get("success"):
        vm = str(r.get("verify_message") or "").strip()
        if vm:
            return f"✅ {vm}"
        return "✅ 已经按照您的要求完成任务。"
    if r.get("task_failed"):
        return f"❌ {r.get('error') or '任务失败'}"
    return f"❌ 执行失败：{r.get('error') or '未知错误'}"


def _format_failed_confirm_tool_result(
    batch: List[Dict[str, Any]],
    tool_context: Optional[Dict[str, Any]],
    result: Dict[str, Any],
) -> str:
    """确认执行失败时与成功分支同风格的 [TOOL_RESULT]，供主助手决策是否改 SQL/查重试。"""
    lines = [
        "[TOOL_RESULT]",
        "action: confirm_failed",
        "status: error",
    ]
    if tool_context:
        a = (tool_context.get("action") or "").strip()
        i = (tool_context.get("intent") or "").strip()
        if a or i:
            lines.append(f"original_intent: action={a} intent={i}")
    err = result.get("error") or "未知错误"
    lines.append(f"error: {err}")
    le = result.get("last_error")
    if le is not None and str(le).strip() and str(le).strip() != str(err).strip():
        lines.append(f"last_error: {le}")
    if "task_failed" in result:
        lines.append(f"task_failed: {bool(result.get('task_failed'))}")
    rr = result.get("recovery_rounds")
    if rr is not None:
        lines.append(f"recovery_rounds: {rr}")
    et = result.get("error_type")
    if et:
        lines.append(f"error_type: {et}")
    msg = result.get("message")
    if msg and str(msg).strip() and str(msg).strip() != str(err).strip():
        lines.append(f"message: {msg}")
    lines.append(f"batch_sql_count: {len(batch)}")
    lines.append("[/TOOL_RESULT]")
    return "\n".join(lines)


def _load_ai_chat_session_rows_for_assistant(session_id: int, limit: int = 80) -> List[Dict[str, Any]]:
    """按时间顺序取最近若干条会话（过久会话只保留尾部以控制 token）。"""
    with db_conn() as conn:
        with conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor) as cur:
            cur.execute(
                """
                SELECT role, content
                FROM (
                    SELECT role, content, id
                    FROM business.ai_chat
                    WHERE session_id = %s
                    ORDER BY id DESC
                    LIMIT %s
                ) sub
                ORDER BY id ASC
                """,
                (session_id, limit),
            )
            rows = cur.fetchall() or []
    return [dict(r) for r in rows]


def _ai_chat_rows_to_llm_messages(rows: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    """与编排器一致：tool 行作为 user 内容传入，便于主模型理解 [TOOL_RESULT]。"""
    out: List[Dict[str, Any]] = []
    for r in rows:
        role = (r.get("role") or "").strip().lower()
        content = r.get("content")
        text = content if isinstance(content, str) else str(content or "")
        if role == "system":
            continue
        if role == "tool":
            out.append({"role": "user", "content": text})
        elif role in ("user", "assistant"):
            out.append({"role": role, "content": text})
    return out


def _format_confirmed_write_tool_result(
    batch: List[Dict[str, Any]],
    results: List[Dict[str, Any]],
    tool_context: Optional[Dict[str, Any]] = None,
) -> str:
    """与 orchestrator 中 TOOL_RESULT 块风格一致，补充 status: confirmed 与执行明细。"""
    lines = [
        "[TOOL_RESULT]",
        "action: confirmed_write",
        "status: confirmed",
    ]
    if tool_context:
        a = (tool_context.get("action") or "").strip()
        i = (tool_context.get("intent") or "").strip()
        if a or i:
            lines.append(f"original_intent: action={a} intent={i}")
    total_affected = 0
    for r in results:
        if not isinstance(r, dict):
            r = {}
        ar = r.get("affected_rows")
        if isinstance(ar, int):
            total_affected += ar
    lines.append(f"affected_rows_total: {total_affected}")
    lines.append(f"batch_count: {len(results)}")
    lines.append("steps:")
    for idx, (item, r) in enumerate(zip(batch, results), start=1):
        if not isinstance(r, dict):
            r = {}
        op = r.get("operation") or ""
        tbl = r.get("table") or ""
        ar = r.get("affected_rows")
        rid = r.get("id")
        lines.append(
            f"  {idx}) operation: {op} | table: {tbl} | affected_rows: {ar} | id: {rid}"
        )
        ret = r.get("returning")
        if ret:
            try:
                tail = json.dumps(
                    _json_safe_confirm_value(ret), ensure_ascii=False, default=str
                )
                if len(tail) > 2000:
                    tail = tail[:2000] + "…"
                lines.append(f"     returning: {tail}")
            except Exception:
                pass
    lines.append("[/TOOL_RESULT]")
    return "\n".join(lines)


def _format_confirmed_write_tool_result_minimal(
    batch: List[Dict[str, Any]],
    results: List[Dict[str, Any]],
    exc_detail: str,
) -> str:
    """详细格式化异常时的兜底块，仍保证 status: confirmed 可被主助手消费。"""
    return "\n".join(
        [
            "[TOOL_RESULT]",
            "action: confirmed_write",
            "status: confirmed",
            "format_error: 详细步骤序列化失败，已降级",
            f"detail: {exc_detail[:1500]}",
            f"batch_sql_count: {len(batch)}",
            f"results_len: {len(results)}",
            "[/TOOL_RESULT]",
        ]
    )


def _normalize_optional_session_id(raw: Any) -> Optional[int]:
    """前端可能传字符串 session_id，统一为 int，非法则视为无会话。"""
    if raw is None:
        return None
    try:
        return int(raw)
    except (TypeError, ValueError):
        return None


def _json_safe_confirm_value(obj: Any) -> Any:
    """将 date/datetime/time 及嵌套结构转为可 JSON 序列化的值（日期时间为 isoformat 字符串）。"""
    if isinstance(obj, datetime):
        return obj.isoformat()
    if isinstance(obj, date):
        return obj.isoformat()
    if isinstance(obj, time):
        return obj.isoformat()
    if isinstance(obj, dict):
        return {k: _json_safe_confirm_value(v) for k, v in obj.items()}
    if isinstance(obj, (list, tuple)):
        return [_json_safe_confirm_value(x) for x in obj]
    return obj


def _returning_dict_from_confirm_results(res_list: Any) -> Dict[str, Any]:
    """从 confirm 的 results 列表合并各步 returning 为单 dict；无则 {}。"""
    out: Dict[str, Any] = {}
    if not isinstance(res_list, list):
        return out
    for r in res_list:
        if not isinstance(r, dict):
            continue
        ret = r.get("returning")
        if isinstance(ret, dict):
            out.update(ret)
    return _json_safe_confirm_value(out)


async def _confirm_persist_task_and_tool_row(
    *,
    sid: Optional[int],
    user_id: Optional[int],
    tool_text: str,
    confirm_success: bool,
    returning_data: Optional[Dict[str, Any]] = None,
    failure_reason: Optional[str] = None,
) -> None:
    """
    更新 ai_tasks，并将 TOOL_RESULT 写入 ai_chat(role=tool)。
    不调用主助手；任务进度由 orchestrator 在专家 done+success 时拼入 [TOOL_RESULT]。
    """
    if sid and _task_pg_pool:
        if confirm_success:
            task = await get_active_task(_task_pg_pool, sid)
            if task:
                step_idx = task["current_step"]
                rd = returning_data if isinstance(returning_data, dict) else {}
                await update_step_result(
                    _task_pg_pool,
                    task["id"],
                    step_idx,
                    "completed",
                    rd,
                )
                if not is_last_step(task):
                    await advance_step(_task_pg_pool, task["id"])
        else:
            task = await get_active_task(_task_pg_pool, sid)
            if task:
                step_idx = task["current_step"]
                fr = (failure_reason or "未知错误").strip()
                await update_step_result(
                    _task_pg_pool,
                    task["id"],
                    step_idx,
                    "failed",
                    {"error": fr},
                )

    if sid is None:
        return

    _insert_ai_chat_row(sid, "tool", tool_text)
    logger.info(
        "ai/confirm: ai_chat role=tool inserted session_id=%s content_len=%s",
        sid,
        len(tool_text or ""),
    )


@app.post("/ai/chat/log-message")
async def ai_chat_log_message(body: ChatLogMessageRequest):
    """将一条消息追加写入 business.ai_chat（如取消写入、前端捕获的确认异常）。"""
    content = (body.content or "").strip()
    if not content:
        raise HTTPException(status_code=400, detail="content 不能为空")
    role = (body.role or "assistant").strip().lower()
    if role not in ("assistant", "user", "tool", "system"):
        raise HTTPException(status_code=400, detail="role 无效")
    try:
        _insert_ai_chat_row(body.session_id, role, content)
    except Exception as e:
        logger.exception("ai/chat/log-message: %s", e)
        raise HTTPException(status_code=500, detail="写入失败") from e
    return {"ok": True}


@app.post("/ai/confirm")
async def ai_confirm(body: WriteConfirmRequest):
    """执行经 preview_sql 审核过的写操作；执行后校验，失败则专家自动修正并重试（最多 MAX_WRITE_CONFIRM_ATTEMPTS 次）。"""

    if body.payload is not None:
        pass
    elif body.frame_id is not None:
        return {
            "success": False,
            "error": "当前版本请使用对话内「写入预览」确认流程，不再支持 frame_id 确认。",
        }
    else:
        raise HTTPException(status_code=400, detail="缺少 payload")

    pl = body.payload if isinstance(body.payload, dict) else {}
    batch = _validate_write_payload_to_batch(pl)
    tc = pl.get("tool_context")
    tool_context = tc if isinstance(tc, dict) else None
    sid = _normalize_optional_session_id(body.session_id)
    result = await _confirm_writes_with_verify_and_recovery(
        batch,
        tool_context,
        user_id=body.user_id,
        session_id=sid,
    )
    if result.get("success"):
        res_list = result.get("results") or []
        try:
            tool_text = _format_confirmed_write_tool_result(
                batch, res_list, tool_context
            )
        except Exception as fe:
            logger.exception(
                "ai/confirm: _format_confirmed_write_tool_result failed: %s", fe
            )
            tool_text = _format_confirmed_write_tool_result_minimal(
                batch, res_list, str(fe)
            )
        try:
            await _confirm_persist_task_and_tool_row(
                sid=sid,
                user_id=body.user_id,
                tool_text=tool_text,
                confirm_success=True,
                returning_data=_returning_dict_from_confirm_results(res_list),
            )
        except Exception as e:
            logger.exception("ai/confirm: TOOL_RESULT 或任务状态落库失败: %s", e)
            err_reply = f"写入已执行，但记录执行结果时出错：{e}"
            result["assistant_reply"] = err_reply
            if sid is not None:
                try:
                    _insert_ai_chat_row(sid, "assistant", err_reply)
                except Exception as e2:
                    logger.warning(
                        "ai/confirm: 写入失败说明到 ai_chat 失败 session_id=%s err=%s",
                        sid,
                        e2,
                    )
    else:
        tool_text = _format_failed_confirm_tool_result(
            batch, tool_context, result
        )
        try:
            await _confirm_persist_task_and_tool_row(
                sid=sid,
                user_id=body.user_id,
                tool_text=tool_text,
                confirm_success=False,
                failure_reason=(result.get("error") or "未知错误"),
            )
        except Exception as e:
            logger.exception(
                "ai/confirm: 失败分支 TOOL_RESULT 或任务状态落库失败: %s", e
            )
            base = (result.get("error") or "执行失败").strip()
            err_reply = f"{base}（记录失败信息时出错：{e}）"
            result["assistant_reply"] = err_reply
            if sid is not None:
                try:
                    _insert_ai_chat_row(sid, "assistant", err_reply)
                except Exception as e2:
                    logger.warning(
                        "ai/confirm: 写入失败说明到 ai_chat 失败 session_id=%s err=%s",
                        sid,
                        e2,
                    )
    return _json_safe_confirm_value(result)


@app.post("/ai/feedback")
async def submit_ai_feedback(body: Dict[str, Any] = Body(...)):
    """用户主动反馈（👎 等）；写入 business.ai_feedback。"""
    raw_uid = body.get("user_id")
    if raw_uid is None:
        raise HTTPException(status_code=400, detail="需要 user_id")
    try:
        user_id_fb = int(raw_uid)
    except (TypeError, ValueError):
        raise HTTPException(status_code=400, detail="user_id 无效")

    sid_raw = body.get("session_id")
    session_id_fb: Optional[int] = None
    if sid_raw is not None:
        try:
            session_id_fb = int(sid_raw)
        except (TypeError, ValueError):
            session_id_fb = None

    chat_id_raw = body.get("chat_id")
    chat_id_fb: Optional[int] = None
    if chat_id_raw is not None:
        try:
            chat_id_fb = int(chat_id_raw)
        except (TypeError, ValueError):
            chat_id_fb = None

    feedback_type = str(body.get("feedback_type") or "thumbs_down").strip() or "thumbs_down"
    content = str(body.get("content") or "")[:2000]

    recent: List[Dict[str, Any]] = []
    if session_id_fb is not None:
        with db_conn() as conn:
            with conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor) as cur:
                cur.execute(
                    """
                    SELECT role, content, created_at::text AS created_at
                    FROM business.ai_chat
                    WHERE session_id = %s
                    ORDER BY id DESC
                    LIMIT 10
                    """,
                    (session_id_fb,),
                )
                rows = list(cur.fetchall() or [])
                recent = [dict(r) for r in reversed(rows)]

    try:
        with db_conn() as conn:
            with conn.cursor() as cur:
                cur.execute(
                    """
                    INSERT INTO business.ai_feedback
                        (session_id, user_id, chat_id, feedback_type, content, recent_messages)
                    VALUES (%s, %s, %s, %s, %s, %s)
                    """,
                    (
                        session_id_fb,
                        user_id_fb,
                        chat_id_fb,
                        feedback_type,
                        content,
                        Json(recent) if recent else None,
                    ),
                )
            conn.commit()
    except Exception as e:
        logger.exception("ai/feedback: %s", e)
        raise HTTPException(status_code=500, detail="反馈失败") from e

    return {"success": True, "message": "感谢反馈，我们会持续改进"}


@app.get("/ai/feedback/summary")
async def feedback_summary(days: int = Query(1, ge=1, le=365)):
    """近 N 天的 incidents / feedback 汇总（运维查看）。"""
    with db_conn() as conn:
        with conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor) as cur:
            cur.execute(
                """
                SELECT stage, COUNT(*)::bigint AS count
                FROM business.ai_incidents
                WHERE created_at >= NOW() - (%s::integer * INTERVAL '1 day')
                GROUP BY stage
                ORDER BY count DESC
                """,
                (days,),
            )
            incidents = [dict(r) for r in cur.fetchall() or []]
            cur.execute(
                """
                SELECT feedback_type, COUNT(*)::bigint AS count
                FROM business.ai_feedback
                WHERE created_at >= NOW() - (%s::integer * INTERVAL '1 day')
                GROUP BY feedback_type
                ORDER BY count DESC
                """,
                (days,),
            )
            feedbacks = [dict(r) for r in cur.fetchall() or []]
            cur.execute(
                """
                SELECT COUNT(DISTINCT session_id)::bigint AS total
                FROM business.ai_chat
                WHERE created_at >= NOW() - (%s::integer * INTERVAL '1 day')
                """,
                (days,),
            )
            row = cur.fetchone()
            total_sessions = int(row["total"]) if row and row.get("total") is not None else 0
    return {
        "period": f"最近 {days} 天",
        "total_sessions": total_sessions,
        "incidents": incidents,
        "feedbacks": feedbacks,
    }


@app.post("/ai/memory/clean-session")
async def api_clean_session(body: Dict[str, Any] = Body(...)):
    """手动触发：清洗指定 session 的对话，提取要点写入 ai_memory。"""
    raw = body.get("session_id")
    if raw is None:
        raise HTTPException(status_code=400, detail="需要 session_id")
    try:
        session_id = int(raw)
    except (TypeError, ValueError):
        raise HTTPException(status_code=400, detail="session_id 必须为整数")
    result = await asyncio.to_thread(clean_session, session_id)
    if not result.get("success"):
        raise HTTPException(status_code=400, detail=result.get("error", "清洗失败"))
    return result


@app.post("/ai/memory/compress-longterm")
async def api_compress_longterm(body: Dict[str, Any] = Body(...)):
    """手动触发：压缩指定用户的长期记忆。"""
    raw = body.get("user_id")
    if raw is None:
        raise HTTPException(status_code=400, detail="需要 user_id")
    try:
        user_id = int(raw)
    except (TypeError, ValueError):
        raise HTTPException(status_code=400, detail="user_id 必须为整数")
    result = await asyncio.to_thread(compress_longterm, user_id)
    if not result.get("success"):
        raise HTTPException(status_code=400, detail=result.get("error", "压缩失败"))
    return result


@app.post("/ai/memory/extract-core")
async def api_extract_core(body: Dict[str, Any] = Body(...)):
    """手动触发：从长期记忆提炼核心经验（scope=core）。"""
    raw = body.get("user_id")
    if raw is None:
        raise HTTPException(status_code=400, detail="需要 user_id")
    try:
        user_id = int(raw)
    except (TypeError, ValueError):
        raise HTTPException(status_code=400, detail="user_id 必须为整数")
    result = await asyncio.to_thread(extract_core_facts, user_id)
    if not result.get("success"):
        raise HTTPException(status_code=400, detail=result.get("error", "提炼失败"))
    return result


@app.get("/ai/memory")
def list_ai_memory(
    user_id: int = Query(..., description="员工用户 id"),
    scope: str = Query(..., description="core、longterm 或 midterm"),
):
    """列出指定用户、指定 scope 的记忆行（不含 session）。"""
    if scope not in ("core", "longterm", "midterm"):
        raise HTTPException(
            status_code=400, detail="scope 须为 core、longterm 或 midterm"
        )
    with db_conn() as conn:
        with conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor) as cur:
            cur.execute(
                """
                SELECT id, user_id, scope, content, created_at, updated_at
                FROM business.ai_memory
                WHERE user_id = %s AND scope = %s
                ORDER BY id ASC
                """,
                (user_id, scope),
            )
            rows = cur.fetchall()
    return {"items": [dict(r) for r in rows]}


class MemoryCreateBody(BaseModel):
    user_id: int
    scope: str
    content: str


@app.post("/ai/memory")
def create_ai_memory(body: MemoryCreateBody):
    if body.scope not in ("core", "longterm"):
        raise HTTPException(
            status_code=400, detail="scope 须为 core 或 longterm"
        )
    c = (body.content or "").strip()
    if not c:
        raise HTTPException(status_code=400, detail="content 不能为空")
    with db_conn() as conn:
        with conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor) as cur:
            cur.execute(
                """
                INSERT INTO business.ai_memory (user_id, scope, content)
                VALUES (%s, %s, %s)
                RETURNING id, user_id, scope, content, created_at, updated_at
                """,
                (body.user_id, body.scope, c),
            )
            row = cur.fetchone()
        conn.commit()
    return dict(row) if row else {}


class MemoryPatchBody(BaseModel):
    content: str
    user_id: int


@app.patch("/ai/memory/{memory_id}")
def patch_ai_memory(memory_id: int, body: MemoryPatchBody):
    c = (body.content or "").strip()
    if not c:
        raise HTTPException(status_code=400, detail="content 不能为空")
    with db_conn() as conn:
        with conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor) as cur:
            cur.execute(
                """
                UPDATE business.ai_memory
                SET content = %s, updated_at = NOW()
                WHERE id = %s AND user_id = %s AND scope IN ('core', 'longterm')
                RETURNING id, user_id, scope, content, created_at, updated_at
                """,
                (c, memory_id, body.user_id),
            )
            row = cur.fetchone()
        conn.commit()
    if not row:
        raise HTTPException(status_code=404, detail="记录不存在或无权限")
    return dict(row)


@app.delete("/ai/memory/{memory_id}")
def delete_ai_memory(
    memory_id: int,
    user_id: int = Query(..., description="须与记录 user_id 一致"),
):
    with db_conn() as conn:
        with conn.cursor() as cur:
            cur.execute(
                """
                DELETE FROM business.ai_memory
                WHERE id = %s AND user_id = %s AND scope IN ('core', 'longterm')
                """,
                (memory_id, user_id),
            )
            n = cur.rowcount
        conn.commit()
    if not n:
        raise HTTPException(status_code=404, detail="记录不存在或无权限")
    return {"ok": True, "id": memory_id}


def _chunk_text_for_sse(text: str, size: int = 48) -> List[str]:
    if not text:
        return [""]
    return [text[i : i + size] for i in range(0, len(text), size)]


def _last_user_message_text(messages: List[ChatMessage]) -> str:
    for m in reversed(messages):
        if m.role == "user":
            return m.content or ""
    return ""


@app.post("/ai/chat")
async def ai_chat(req: ChatRequest, request: Request):
    if req.pending_token:
        raise HTTPException(
            status_code=400,
            detail="当前版本不支持 pendingToken 恢复，请发起新对话。",
        )
    if req.confirm:
        raise HTTPException(
            status_code=400,
            detail="当前版本不支持在此确认数据库操作；请使用旧后端或关闭「确认」类流程。",
        )

    session_id: Optional[int] = req.session_id
    user_id: Optional[int] = req.user_id

    if user_id is not None and session_id is None:
        with db_conn() as conn:
            with conn.cursor() as cur:
                cur.execute(
                    """
                    INSERT INTO business.ai_sessions (user_id, status)
                    VALUES (%s, 'in_progress')
                    RETURNING id
                    """,
                    (user_id,),
                )
                row = cur.fetchone()
                session_id = int(row[0]) if row else None
            conn.commit()

    history_messages = [{"role": m.role, "content": m.content} for m in req.messages]
    file_summaries: List[str] = []

    if req.files:
        for file_payload in req.files:
            parsed = await asyncio.to_thread(parse_file, file_payload)
            if parsed.get("summary"):
                file_summaries.append(str(parsed["summary"]))
            p_type = parsed.get("type")
            if p_type == "text":
                extra = (parsed.get("content") or "").strip()
                if extra:
                    for m in reversed(history_messages):
                        if m.get("role") == "user":
                            m["content"] = (m.get("content") or "") + "\n\n【附件】\n" + extra
                            break

    if user_id is not None and session_id is not None:
        mem_auto = await asyncio.to_thread(
            maybe_clean_previous_session, user_id, session_id
        )
        if mem_auto.get("auto_cleaned"):
            logger.info(
                "memory: 已自动整理上一会话并写入 ai_memory | user_id=%s previous_session_id=%s",
                user_id,
                mem_auto.get("previous_session_id"),
            )
        elif mem_auto.get("error"):
            logger.warning(
                "memory: 自动整理上一会话失败 | user_id=%s err=%s",
                user_id,
                mem_auto.get("error"),
            )

    system_prompt = build_prompt("assistant", user_id=user_id)
    api_messages: List[Dict[str, Any]] = [
        {"role": "system", "content": system_prompt},
    ] + history_messages

    use_expert = bool(req.invite_expert)
    user_msg_for_log = _last_user_message_text(req.messages)

    # ── SSE 流式（前端默认 stream=true）──────────────────────────────
    if req.stream:

        async def sse_gen():
            try:
                yield _sse_line({"type": "start", "session_id": session_id})
                meta_payload = {
                    "type": "meta",
                    "intent": "chat",
                    "requiresConfirm": False,
                    "pendingToken": None,
                    "fileSummaries": file_summaries,
                    "data": {"main": [], "sub": []},
                }
                yield _sse_line(meta_payload)

                work_msgs: List[Dict[str, Any]] = list(api_messages)

                if user_id is not None and session_id is not None:
                    _insert_ai_chat_row(session_id, "user", user_msg_for_log)

                recent_tool_blocks: List[tuple[str, str]] = []

                for _round in range(MAX_ORCHESTRATOR_ROUNDS):
                    if await request.is_disconnected():
                        if user_id is not None and session_id is not None:
                            _insert_ai_chat_row(
                                session_id, "system", _ORCH_SYSTEM_ABORT
                            )
                        break
                    logger.info(
                        "orchestrator: 进入第 %d/%d 轮",
                        _round + 1,
                        MAX_ORCHESTRATOR_ROUNDS,
                    )
                    full_reply = ""
                    last_vis_len = 0
                    if use_expert:
                        full_reply = await asyncio.to_thread(chat_expert_only, work_msgs)
                        vis = _user_visible_stream_prefix(full_reply)
                        for part in _chunk_text_for_sse(vis):
                            yield _sse_line({"type": "delta", "text": part})
                    else:
                        async for piece in chat_stream(work_msgs):
                            if piece:
                                full_reply += piece
                                vis = _user_visible_stream_prefix(full_reply)
                                if len(vis) > last_vis_len:
                                    yield _sse_line(
                                        {"type": "delta", "text": vis[last_vis_len:]}
                                    )
                                    last_vis_len = len(vis)

                    # ── V9：检测 [PLAN] 块，写入任务状态机 ──────────────────────
                    if "[PLAN]" in full_reply:
                        plan_intent, plan_steps = parse_plan(full_reply)
                        if (
                            plan_intent
                            and plan_steps
                            and session_id
                            and user_id is not None
                            and _task_pg_pool is not None
                        ):
                            await create_task(
                                pool=_task_pg_pool,
                                session_id=session_id,
                                user_id=user_id,
                                original_intent=plan_intent,
                                steps=plan_steps,
                            )
                        full_reply = re.sub(
                            r"\[PLAN\].*?\[/PLAN\]", "", full_reply, flags=re.DOTALL
                        ).strip()
                        if "[PLAN]" in full_reply:
                            _pi = full_reply.find("[PLAN]")
                            full_reply = full_reply[:_pi].rstrip()
                    # ── V9 结束 ──────────────────────────────────────────────────

                    logged_reply = _strip_assistant_fabrication_after_tool(full_reply)
                    assistant_user_visible = _user_visible_stream_prefix(full_reply).rstrip()
                    if user_id is not None and session_id is not None:
                        _insert_ai_chat_row(session_id, "assistant", assistant_user_visible)

                    tool_block = _extract_tool_block(full_reply)
                    if not tool_block:
                        break

                    current_ai = (
                        (tool_block.get("action") or "").strip(),
                        (tool_block.get("intent") or "").strip(),
                    )
                    recent_tool_blocks.append(current_ai)
                    if len(recent_tool_blocks) >= 3:
                        t3 = recent_tool_blocks[-3:]
                        if t3[0] == t3[1] == t3[2]:
                            log_incident(
                                stage="expert",
                                error_message=_ORCH_MSG_DUP_TOOL,
                                session_id=session_id,
                                user_id=user_id,
                                user_input=user_msg_for_log,
                                context={
                                    "spin": True,
                                    "action": current_ai[0],
                                    "intent": current_ai[1],
                                },
                            )
                            yield _sse_line(
                                {
                                    "type": "delta",
                                    "text": f"\n\n{_ORCH_MSG_DUP_TOOL}\n\n",
                                }
                            )
                            if user_id is not None and session_id is not None:
                                _insert_ai_chat_row(
                                    session_id, "assistant", _ORCH_MSG_DUP_TOOL
                                )
                            break

                    yield _sse_line({"type": "delta", "text": _SSE_TOOL_PLEASE_WAIT})

                    result = await _call_tool_expert(
                        tool_block, user_id, session_id=session_id
                    )

                    if _is_write_preview_result(result):
                        result_text = _format_tool_result(tool_block, result)
                        if user_id is not None and session_id is not None:
                            _insert_ai_chat_row(session_id, "tool", result_text)
                        _wp_log = _write_preview_user_visible_log(result)
                        # delta 仍只下发简短提示，避免与前端卡片重复占屏；完整摘要写入 ai_chat，并由前端在 write_preview 事件中并入 content
                        yield _sse_line(
                            {
                                "type": "delta",
                                "text": f"\n\n{_WRITE_PREVIEW_CONFIRM_HINT}\n\n",
                            }
                        )
                        if user_id is not None and session_id is not None:
                            _insert_ai_chat_row(session_id, "assistant", _wp_log)
                        payload_wp = result.get("payload") or {
                            "sql": result.get("sql"),
                            "params": result.get("params") or [],
                        }
                        if isinstance(payload_wp, dict):
                            payload_wp = _merge_tool_context_into_payload(
                                payload_wp, tool_block
                            )
                        yield _sse_line(
                            {
                                "type": "write_preview",
                                "stage": result.get("stage"),
                                "summary": result.get("summary"),
                                "expert_explanation": result.get("expert_explanation"),
                                "highlight_words": result.get("highlight_words") or [],
                                "payload": payload_wp,
                                "operation": result.get("operation"),
                                "table": result.get("table"),
                            }
                        )
                        break

                    result_text = await _format_tool_result_with_task_progress(
                        tool_block, result, session_id, user_id
                    )
                    if user_id is not None and session_id is not None:
                        _insert_ai_chat_row(session_id, "tool", result_text)

                    work_msgs.append({"role": "assistant", "content": logged_reply})
                    work_msgs.append({"role": "user", "content": result_text})

                else:
                    log_incident(
                        stage="expert",
                        error_message=_ORCH_MSG_MAX_ROUNDS,
                        session_id=session_id,
                        user_id=user_id,
                        user_input=user_msg_for_log,
                        context={"max_orchestrator_rounds": MAX_ORCHESTRATOR_ROUNDS},
                    )
                    yield _sse_line(
                        {
                            "type": "delta",
                            "text": f"\n\n{_ORCH_MSG_MAX_ROUNDS}\n\n",
                        }
                    )
                    if user_id is not None and session_id is not None:
                        _insert_ai_chat_row(
                            session_id, "assistant", _ORCH_MSG_MAX_ROUNDS
                        )

                yield _sse_line(
                    {
                        "type": "done",
                        "data": {"main": [], "sub": []},
                    }
                )
            except Exception as e:
                logger.exception("ai/chat SSE: %s", e)
                yield _sse_line({"type": "error", "detail": str(e)})
                yield _sse_line({"type": "done", "data": {"main": [], "sub": []}})

        return StreamingResponse(
            sse_gen(),
            media_type="text/event-stream",
            headers={
                "Cache-Control": "no-cache",
                "Connection": "keep-alive",
                "X-Accel-Buffering": "no",
            },
        )

    # ── 非流式 JSON ─────────────────────────────────────────────────
    work_msgs = list(api_messages)
    pr: Optional[Dict[str, Any]] = None
    pr_tool_block: Optional[Dict[str, str]] = None
    reply = ""

    try:
        if user_id is not None and session_id is not None:
            _insert_ai_chat_row(session_id, "user", user_msg_for_log)

        recent_tool_blocks_ns: List[tuple[str, str]] = []

        for _round in range(MAX_ORCHESTRATOR_ROUNDS):
            if await request.is_disconnected():
                if user_id is not None and session_id is not None:
                    _insert_ai_chat_row(session_id, "system", _ORCH_SYSTEM_ABORT)
                break
            logger.info(
                "orchestrator: 进入第 %d/%d 轮",
                _round + 1,
                MAX_ORCHESTRATOR_ROUNDS,
            )
            if use_expert:
                full_reply = await asyncio.to_thread(chat_expert_only, work_msgs)
            else:
                full_reply = await asyncio.to_thread(chat, work_msgs)

            # ── V9：检测 [PLAN] 块，写入任务状态机 ──────────────────────
            if "[PLAN]" in full_reply:
                plan_intent, plan_steps = parse_plan(full_reply)
                if (
                    plan_intent
                    and plan_steps
                    and session_id
                    and user_id is not None
                    and _task_pg_pool is not None
                ):
                    await create_task(
                        pool=_task_pg_pool,
                        session_id=session_id,
                        user_id=user_id,
                        original_intent=plan_intent,
                        steps=plan_steps,
                    )
                full_reply = re.sub(
                    r"\[PLAN\].*?\[/PLAN\]", "", full_reply, flags=re.DOTALL
                ).strip()
                if "[PLAN]" in full_reply:
                    _pi = full_reply.find("[PLAN]")
                    full_reply = full_reply[:_pi].rstrip()
            # ── V9 结束 ──────────────────────────────────────────────────

            logged_reply = _strip_assistant_fabrication_after_tool(full_reply)
            assistant_user_visible = _user_visible_stream_prefix(full_reply).rstrip()
            if user_id is not None and session_id is not None:
                _insert_ai_chat_row(session_id, "assistant", assistant_user_visible)

            reply = assistant_user_visible
            tool_block = _extract_tool_block(full_reply)
            if not tool_block:
                break

            cur_ai = (
                (tool_block.get("action") or "").strip(),
                (tool_block.get("intent") or "").strip(),
            )
            recent_tool_blocks_ns.append(cur_ai)
            if len(recent_tool_blocks_ns) >= 3:
                t3 = recent_tool_blocks_ns[-3:]
                if t3[0] == t3[1] == t3[2]:
                    log_incident(
                        stage="expert",
                        error_message=_ORCH_MSG_DUP_TOOL,
                        session_id=session_id,
                        user_id=user_id,
                        user_input=user_msg_for_log,
                        context={
                            "spin": True,
                            "action": cur_ai[0],
                            "intent": cur_ai[1],
                        },
                    )
                    reply = assistant_user_visible + "\n\n" + _ORCH_MSG_DUP_TOOL
                    if user_id is not None and session_id is not None:
                        _insert_ai_chat_row(
                            session_id, "assistant", _ORCH_MSG_DUP_TOOL
                        )
                    break

            result = await _call_tool_expert(
                tool_block, user_id, session_id=session_id
            )

            if _is_write_preview_result(result):
                result_text = _format_tool_result(tool_block, result)
                _wp_log = _write_preview_user_visible_log(result)
                if user_id is not None and session_id is not None:
                    _insert_ai_chat_row(session_id, "tool", result_text)
                    _insert_ai_chat_row(session_id, "assistant", _wp_log)
                reply = assistant_user_visible + _SSE_TOOL_PLEASE_WAIT + _wp_log
                pr = result
                pr_tool_block = tool_block
                break

            result_text = await _format_tool_result_with_task_progress(
                tool_block, result, session_id, user_id
            )
            if user_id is not None and session_id is not None:
                _insert_ai_chat_row(session_id, "tool", result_text)

            work_msgs.append({"role": "assistant", "content": logged_reply})
            work_msgs.append({"role": "user", "content": result_text})
        else:
            log_incident(
                stage="expert",
                error_message=_ORCH_MSG_MAX_ROUNDS,
                session_id=session_id,
                user_id=user_id,
                user_input=user_msg_for_log,
                context={"max_orchestrator_rounds": MAX_ORCHESTRATOR_ROUNDS},
            )
            sep = "\n\n" if reply else ""
            reply = (reply or "") + sep + _ORCH_MSG_MAX_ROUNDS
            if user_id is not None and session_id is not None:
                _insert_ai_chat_row(session_id, "assistant", _ORCH_MSG_MAX_ROUNDS)
    except Exception as e:
        raise HTTPException(status_code=502, detail=str(e)) from e

    out: Dict[str, Any] = {
        "reply": reply,
        "content": reply,
        "data": {"main": [], "sub": []},
        "intent": "chat",
        "requiresConfirm": False,
        "pendingToken": None,
        "fileSummaries": file_summaries,
        "confirm_required": False,
    }
    if session_id is not None:
        out["session_id"] = session_id
    if pr and pr.get("success"):
        out["stage"] = "preview"
        out["summary"] = pr.get("summary")
        out["expert_explanation"] = pr.get("expert_explanation")
        out["highlight_words"] = pr.get("highlight_words") or []
        base_pl = pr.get("payload") or {
            "sql": pr.get("sql"),
            "params": pr.get("params") or [],
        }
        if isinstance(base_pl, dict) and pr_tool_block:
            out["payload"] = _merge_tool_context_into_payload(base_pl, pr_tool_block)
        else:
            out["payload"] = base_pl
        out["operation"] = pr.get("operation")
        out["table"] = pr.get("table")
    return out
