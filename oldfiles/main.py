from fastapi import FastAPI, HTTPException, Body, Query
from fastapi.encoders import jsonable_encoder
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import Response, StreamingResponse
from pydantic import BaseModel, ConfigDict, Field
from typing import Any, Dict, List, Optional, Tuple
from contextlib import asynccontextmanager
import re
import io
import json
import uuid
from datetime import date, datetime, time, timezone
from decimal import Decimal
import psycopg2
import psycopg2.extras
from psycopg2.extras import Json
import os
import asyncio
import logging
import sys
import unicodedata
import httpx
from pathlib import Path

from tools.database import close_all_pools, db_conn

logger = logging.getLogger(__name__)


def _env_bool(key: str, default: bool) -> bool:
    v = os.environ.get(key)
    if v is None:
        return default
    return v.strip().lower() in ("1", "true", "yes", "on")


def _cors_allow_origins_and_credentials() -> tuple[list[str], bool]:
    """
    浏览器禁止 Access-Control-Allow-Origin: * 与 Access-Control-Allow-Credentials: true 同时出现。
    - CORS_ORIGINS 未设置：默认常见本地前端端口白名单，credentials 默认可为 true。
    - CORS_ORIGINS=*：任意源时 credentials 强制为 false。
    - 其它：逗号分隔白名单；credentials 由 CORS_ALLOW_CREDENTIALS 控制（默认 true）。
    """
    raw = (os.environ.get("CORS_ORIGINS") or "").strip()

    def _warn_if_star_overrides_credentials() -> None:
        if _env_bool("CORS_ALLOW_CREDENTIALS", False):
            logger.warning(
                "CORS_ALLOW_CREDENTIALS=true 与 CORS_ORIGINS=* 不兼容，已强制设为 false"
            )

    if raw == "*":
        _warn_if_star_overrides_credentials()
        return ["*"], False
    if not raw:
        origins = [
            "http://127.0.0.1:5173",
            "http://localhost:5173",
            "http://127.0.0.1:3000",
            "http://localhost:3000",
        ]
        return origins, _env_bool("CORS_ALLOW_CREDENTIALS", True)
    origins = [p.strip() for p in raw.split(",") if p.strip()]
    if not origins:
        return (
            [
                "http://127.0.0.1:5173",
                "http://localhost:5173",
            ],
            _env_bool("CORS_ALLOW_CREDENTIALS", True),
        )
    if origins == ["*"]:
        _warn_if_star_overrides_credentials()
        return ["*"], False
    return origins, _env_bool("CORS_ALLOW_CREDENTIALS", True)


@asynccontextmanager
async def _app_lifespan(app: FastAPI):
    # 保证 engine.collecting 的 INFO（如完整 system prompt）在 uvicorn 控制台可见
    _fmt = logging.Formatter("%(levelname)s %(name)s: %(message)s")
    _cl = logging.getLogger("engine.collecting")
    _cl.setLevel(logging.INFO)
    if not _cl.handlers:
        _h = logging.StreamHandler(sys.stderr)
        _h.setLevel(logging.INFO)
        _h.setFormatter(_fmt)
        _cl.addHandler(_h)
    _cl.propagate = False
    for _name in (
        "__main__",
        "main",
        "engine.workflow_engine_exec",
        "engine.pending_submit_v15b",
    ):
        _lg = logging.getLogger(_name)
        _lg.setLevel(logging.INFO)
        if not _lg.handlers:
            _hh = logging.StreamHandler(sys.stderr)
            _hh.setLevel(logging.INFO)
            _hh.setFormatter(_fmt)
            _lg.addHandler(_hh)
        _lg.propagate = False
    yield
    close_all_pools()


app = FastAPI(title="AI·ERP Backend", lifespan=_app_lifespan)
print(
    f"[aierp] 已加载 main.py（请确认 uvicorn 工作目录与此一致）: {Path(__file__).resolve()}",
    flush=True,
)

# ── CORS（* 与 credentials 互斥，见 _cors_allow_origins_and_credentials）────────
_cors_origins, _cors_cred = _cors_allow_origins_and_credentials()
print(
    f"[aierp] CORS allow_origins={_cors_origins!r} allow_credentials={_cors_cred}",
    flush=True,
)
app.add_middleware(
    CORSMiddleware,
    allow_origins=_cors_origins,
    allow_credentials=_cors_cred,
    allow_methods=["*"],
    allow_headers=["*"],
)

IDENT_RE = re.compile(r"^[a-zA-Z_][a-zA-Z0-9_]*$")
ALLOWED_SCHEMAS = frozenset({"meta", "business", "ai_assistant"})

# ── 意图启发式：先 NFKC，再遮蔽「否定 + 删除/DROP」类子串，最后匹配下列正则，减少误报 ──
# 否定删除（中英）：「我不想删除这条记录」等不再触发删数据/DELETE 类规则（主语可在否定词前）
_NEGATED_DELETE_SPAN = re.compile(
    r"(?is)"
    r"(?:"
    r".{0,16}?"
    r"(?:不想|不要|不能|请勿|无需|不必|避免|取消|禁止|决不|绝不|千万别|切莫)"
    r".{0,28}?"
    r"(?:删表|删除表|删除|删掉|清空\s*表|DROP\s+TABLE|DELETE\s+FROM|\bDELETE\b|TRUNCATE)"
    r"|"
    r".{0,12}?"
    r"(?:不|别|勿|没|没有|莫|未|甭)"
    r"(?:是|想|要|会|能|该|应|宜|必|需|准备|打算)?"
    r".{0,16}?"
    r"(?:删表|删除表|删除|删掉)"
    r"|"
    r"(?:(?:don\x27t|do\s+not|never|won\x27t|will\s+not|did\s+not)\s+.{0,48}?"
    r"(?:\bDELETE\b|DROP\s+TABLE|TRUNCATE|DELETE\s+FROM))"
    r")"
)


def _prepare_user_text_for_intent_regex(raw: str) -> str:
    """NFKC（全角英文/数字→半角）+ 遮蔽否定删除子串，再供工具调用断言等正则使用。"""
    t = unicodedata.normalize("NFKC", raw or "")
    return _NEGATED_DELETE_SPAN.sub(lambda m: " " * len(m.group(0)), t)


# 用户话里明显涉及读库/写库/SQL 时，用于检测模型是否应调用工具（见 _assert_tool_was_called）
_FORCE_QUERY_INTENT = re.compile(
    r"(?is)(建表|创建表|新建表|删表|删除表|改表|修改表|表结构|增删改)"
    r"|(创建.{0,48}?表|新建.{0,24}?表|建立.{0,48}?表)"
    r"|(CREATE\s+TABLE|DROP\s+TABLE|ALTER\s+TABLE|TRUNCATE\s+TABLE)"
    r"|(\b(INSERT|UPDATE|DELETE|MERGE|SELECT)\b)"
    r"|(查询|统计|列出|查看).{0,12}?(表|数据|记录|字段|结构)",
)


def _assert_tool_was_called(response_message: Any, user_text: str) -> None:
    """如果用户明显在要求数据库操作，但模型没调工具，拒绝该回复。"""
    prepared = _prepare_user_text_for_intent_regex(user_text)
    if not _FORCE_QUERY_INTENT.search(prepared):
        return
    has_tool_call = False
    if isinstance(response_message, dict):
        if response_message.get("tool_calls"):
            has_tool_call = True
        else:
            for b in response_message.get("content") or []:
                if not isinstance(b, dict):
                    continue
                if b.get("type") in ("tool_use", "function"):
                    has_tool_call = True
                    break
    else:
        if getattr(response_message, "tool_calls", None):
            has_tool_call = True
    if not has_tool_call:
        raise ValueError("模型未调用工具，拒绝接受该回复")

# 设计期通用 CRUD：仍禁止向量与密码哈希直写
BLOCK_COLUMNS = frozenset({"password_hash", "embedding", "intent_embedding"})


def assert_ident(name: str, label: str = "identifier") -> None:
    if not name or not IDENT_RE.match(name):
        raise HTTPException(status_code=400, detail=f"Invalid {label}: {name}")


def fetch_column_meta(cur, schema: str, table: str) -> List[Dict[str, Any]]:
    cur.execute(
        """
        SELECT c.column_name AS name,
               c.data_type AS data_type,
               c.udt_name AS udt_name,
               c.is_nullable AS is_nullable,
               c.column_default AS column_default,
               c.is_generated AS is_generated,
               c.identity_generation AS identity_generation
        FROM information_schema.columns c
        WHERE c.table_schema = %s AND c.table_name = %s
        ORDER BY c.ordinal_position
        """,
        (schema, table),
    )
    return [dict(r) for r in cur.fetchall()]


def writable_columns(meta: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    out = []
    for r in meta:
        name = r["name"]
        if name in BLOCK_COLUMNS:
            continue
        if r.get("udt_name") == "vector":
            continue
        if (r.get("is_generated") or "NEVER") == "ALWAYS":
            continue
        out.append(r)
    return out


def writable_names(meta: List[Dict[str, Any]]) -> List[str]:
    return [r["name"] for r in writable_columns(meta)]


def coerce_cell(val: Any, udt_name: str, data_type: str, nullable: str) -> Any:
    if val is None:
        return None
    if val == "" and nullable == "YES":
        return None
    if isinstance(val, (dict, list)):
        return Json(val)
    return val


def table_has_id(meta: List[Dict[str, Any]]) -> bool:
    return any(r["name"] == "id" for r in meta)


LIKE_ESC = "!"  # PostgreSQL LIKE/ILIKE 转义符，避免与反斜杠混淆


def filterable_column_names(meta: List[Dict[str, Any]]) -> List[str]:
    out: List[str] = []
    for r in meta:
        n = r["name"]
        if n in BLOCK_COLUMNS:
            continue
        if r.get("udt_name") == "vector":
            continue
        out.append(n)
    return out


def _like_escape(s: str, esc: str = LIKE_ESC) -> str:
    if len(esc) != 1:
        raise ValueError("escape must be one character")
    return (
        s.replace(esc, esc + esc).replace("%", esc + "%").replace("_", esc + "_")
    )


def parse_filters_query(filters_raw: Optional[str]) -> Dict[str, str]:
    if not filters_raw or not str(filters_raw).strip():
        return {}
    try:
        o = json.loads(filters_raw)
    except json.JSONDecodeError:
        raise HTTPException(status_code=400, detail="filters 须为合法 JSON 对象")
    if not isinstance(o, dict):
        raise HTTPException(status_code=400, detail="filters 须为 JSON 对象")
    out: Dict[str, str] = {}
    for k, v in o.items():
        if not isinstance(k, str):
            continue
        vs = str(v).strip() if v is not None else ""
        if vs:
            out[k] = vs
    return out


def build_fuzzy_where_clause(
    meta: List[Dict[str, Any]],
    filters: Dict[str, str],
    q_global: Optional[str],
) -> tuple[str, List[Any]]:
    """生成 WHERE 之后的片段（不含 WHERE）及参数；多列 filters 为 AND，q 为整行 OR。"""
    allowed = set(filterable_column_names(meta))
    parts: List[str] = []
    params: List[Any] = []
    esc = LIKE_ESC

    for col, val in filters.items():
        if col not in allowed or not IDENT_RE.match(col):
            continue
        pat = "%" + _like_escape(val, esc) + "%"
        parts.append(f'"{col}"::text ILIKE %s ESCAPE \'{esc}\'')
        params.append(pat)

    qt = (q_global or "").strip()
    if qt:
        pat = "%" + _like_escape(qt, esc) + "%"
        ors: List[str] = []
        for col in sorted(allowed):
            if not IDENT_RE.match(col):
                continue
            ors.append(f'"{col}"::text ILIKE %s ESCAPE \'{esc}\'')
            params.append(pat)
        if ors:
            parts.append("(" + " OR ".join(ors) + ")")

    if not parts:
        return "", []
    return " AND ".join(parts), params


def _order_expr_for_column(meta_row: Dict[str, Any]) -> str:
    col = meta_row["name"]
    if meta_row.get("udt_name") == "vector":
        return f'"{col}"::text'
    return f'"{col}"'


def build_order_clause(
    meta: List[Dict[str, Any]],
    sort_by: Optional[str],
    sort_order_raw: Optional[str],
) -> str:
    """全表排序在 LIMIT/OFFSET 之前执行（由 SQL 执行顺序保证）。"""
    by_name = {r["name"]: r for r in meta}
    names_order = [r["name"] for r in meta]
    order = (sort_order_raw or "asc").lower()
    if order not in ("asc", "desc"):
        order = "asc"
    col = sort_by
    if (
        not col
        or col not in by_name
        or not IDENT_RE.match(col)
    ):
        if "id" in by_name:
            col = "id"
            order = "desc"
        elif names_order:
            col = names_order[0]
            order = "asc"
        else:
            return "ORDER BY 1 ASC NULLS LAST"
    expr = _order_expr_for_column(by_name[col])
    return f"ORDER BY {expr} {order.upper()} NULLS LAST"


def normalize_sort_order_param(raw: Optional[str]) -> Optional[str]:
    if not raw or not str(raw).strip():
        return None
    x = str(raw).strip().lower()
    return x if x in ("asc", "desc") else None


def fetch_table_page(
    cur: Any,
    schema: str,
    table_name: str,
    meta: List[Dict[str, Any]],
    filters_raw: Optional[str],
    q: Optional[str],
    sort_by: Optional[str],
    sort_order: Optional[str],
    limit: int,
    offset: int,
    *,
    with_total: bool,
) -> tuple[List[Dict[str, Any]], Optional[int]]:
    fd = (
        parse_filters_query(filters_raw)
        if filters_raw and str(filters_raw).strip()
        else {}
    )
    where_sql, wparams = build_fuzzy_where_clause(meta, fd, q)
    order_clause = build_order_clause(meta, sort_by, sort_order)
    base_from = f'FROM "{schema}"."{table_name}"'
    frag = (" WHERE " + where_sql) if where_sql else ""

    total: Optional[int] = None
    if with_total:
        cur.execute(f"SELECT COUNT(*) AS _cnt {base_from}{frag}", tuple(wparams))
        total = int(cur.fetchone()["_cnt"])

    cur.execute(
        f"SELECT * {base_from}{frag} {order_clause} LIMIT %s OFFSET %s",
        (*wparams, limit, offset),
    )
    rows = [dict(r) for r in cur.fetchall()]
    return rows, total


# ── Routes ──────────────────────────────────────────────────────────────────
@app.get("/")
def root():
    return {"status": "ok", "project": "AI·ERP"}

@app.get("/tables")
def list_tables():
    with db_conn() as conn:
        with conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor) as cur:
            cur.execute("""
                SELECT table_schema AS schema, table_name AS table
                FROM information_schema.tables
                WHERE table_schema IN ('meta', 'business', 'ai_assistant')
                  AND table_type = 'BASE TABLE'
                ORDER BY table_schema, table_name
            """)
            rows = cur.fetchall()
        return [dict(r) for r in rows]

def _schema_by_table_name(conn) -> dict:
    with conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor) as cur:
        cur.execute(
            """
            SELECT table_name, table_schema
            FROM information_schema.tables
            WHERE table_schema IN ('meta', 'business', 'ai_assistant') AND table_type = 'BASE TABLE'
            """
        )
        m = {}
        for r in cur.fetchall():
            m[r["table_name"]] = r["table_schema"]
        return m


@app.get("/nav/tables")
def list_nav_tables():
    """按 meta.meta_tables.nav_section 分组，供右侧三大块导航。"""
    with db_conn() as conn:
        schema_map = _schema_by_table_name(conn)
        with conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor) as cur:
            cur.execute(
                """
                SELECT table_name, table_cn, nav_section, nav_order, nav_parent_table
                FROM meta.meta_tables
                WHERE nav_section IS NOT NULL
                ORDER BY nav_section, nav_order, table_name
                """
            )
            rows = cur.fetchall()
        out = {"design": [], "prep": [], "run": []}
        for r in rows:
            sec = r.get("nav_section")
            if sec not in out:
                continue
            sch = schema_map.get(r["table_name"])
            if not sch:
                continue
            out[sec].append(
                {
                    "schema": sch,
                    "table_name": r["table_name"],
                    "table_cn": r["table_cn"],
                    "nav_order": r["nav_order"],
                    "nav_parent_table": r["nav_parent_table"],
                }
            )
        return out


@app.get("/table-columns/{table_name}")
def get_table_column_labels(table_name: str):
    """列展示名：优先 menu_label_cn（阶段2 只读表头）。"""
    with db_conn() as conn:
        with conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor) as cur:
            cur.execute(
                """
                SELECT field_name, COALESCE(menu_label_cn, field_name) AS menu_label_cn,
                       COALESCE(menu_label_en, field_name) AS menu_label_en,
                       data_type
                FROM meta.meta_fields
                WHERE table_name = %s
                ORDER BY id
                """,
                (table_name,),
            )
            cols = cur.fetchall()
        return {"columns": [dict(c) for c in cols]}


@app.get("/table/{schema}/{table_name}")
def get_table(
    schema: str,
    table_name: str,
    limit: int = 500,
    offset: int = Query(0, ge=0, le=1_000_000),
    sort_by: Optional[str] = Query(None, description="排序列名，须为表内列"),
    sort_order: Optional[str] = Query(
        None, description="asc 或 desc；缺省时由服务端默认（有 id 则 id desc）"
    ),
    filters: Optional[str] = Query(
        None,
        description='JSON：{"列名":"关键词"}，多列 AND，每列 ILIKE 模糊',
    ),
    q: Optional[str] = Query(
        None,
        description="任意可筛选列 OR 模糊；与 filters 同时存在时为 AND",
    ),
):
    if schema not in ALLOWED_SCHEMAS:
        raise HTTPException(status_code=400, detail="Invalid schema")
    assert_ident(table_name, "table_name")
    lim = min(max(1, limit), 5000)
    off = max(0, min(offset, 1_000_000))
    so = normalize_sort_order_param(sort_order)
    with db_conn() as conn:
        with conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor) as cur:
            try:
                meta = fetch_column_meta(cur, schema, table_name)
                if not meta:
                    raise HTTPException(status_code=404, detail="Table not found")
                rows, total = fetch_table_page(
                    cur,
                    schema,
                    table_name,
                    meta,
                    filters,
                    q,
                    sort_by,
                    so,
                    lim,
                    off,
                    with_total=True,
                )
            except HTTPException:
                raise
            except Exception as e:
                raise HTTPException(status_code=400, detail=str(e)) from e
            return {"rows": rows, "total": total}


def _export_cell_value(val: Any, col_meta: Optional[Dict[str, Any]]) -> Any:
    """写入 xlsx 的单元格值：避免向量/二进制；日期与 Decimal 可识别。
    PostgreSQL uuid 列经 psycopg2 为 uuid.UUID，openpyxl 不能直接写入，须转 str。"""
    if val is None:
        return None
    udt = (col_meta or {}).get("udt_name")
    if udt == "vector":
        return "(向量，已省略)"
    if isinstance(val, uuid.UUID):
        return str(val)
    if isinstance(val, (bytes, memoryview)):
        return "(二进制)"
    if isinstance(val, (dict, list)):
        return json.dumps(val, ensure_ascii=False)
    if isinstance(val, Decimal):
        return str(val)
    # openpyxl / Excel 不接受带 tzinfo 的 datetime（PostgreSQL timestamptz 常见）
    if isinstance(val, datetime):
        if val.tzinfo is not None:
            return val.astimezone(timezone.utc).replace(tzinfo=None)
        return val
    if isinstance(val, date):
        return val
    if isinstance(val, time):
        if val.tzinfo is not None:
            return val.replace(tzinfo=None)
        return val
    if isinstance(val, bool):
        return val
    if isinstance(val, (int, float, str)):
        return val
    # 其余类型（如 INET 等适配器对象）统一转字符串，避免 openpyxl 抛错导致 500
    return str(val)


def _resolve_export_columns(
    meta: List[Dict[str, Any]], cols_param: Optional[str]
) -> List[str]:
    all_names = [r["name"] for r in meta]
    name_set = set(all_names)
    if not cols_param or not cols_param.strip():
        return all_names
    out: List[str] = []
    seen = set()
    for raw in cols_param.split(","):
        c = raw.strip()
        if not c or c in seen:
            continue
        if not IDENT_RE.match(c) or c not in name_set:
            continue
        seen.add(c)
        out.append(c)
    return out if out else all_names


@app.get("/export/xlsx/{schema}/{table_name}")
def export_table_xlsx(
    schema: str,
    table_name: str,
    limit: int = 5000,
    cols: Optional[str] = None,
    filters: Optional[str] = Query(None, description="与列表查询相同的 filters JSON"),
    q: Optional[str] = Query(None, description="与列表查询相同的任意列模糊"),
    sort_by: Optional[str] = Query(None, description="与列表相同的排序列"),
    sort_order: Optional[str] = Query(None, description="asc 或 desc"),
):
    """
    导出为简单 .xlsx（Office 2007+ / WPS），仅数据表与文本单元格，无宏、无扩展样式。
    cols：逗号分隔字段名，顺序与界面一致；不传则按库表列顺序。
    """
    try:
        from openpyxl import Workbook
    except ImportError as e:
        raise HTTPException(
            status_code=503,
            detail="缺少 openpyxl，无法导出。请在运行后端的 Python 环境中执行: pip install openpyxl",
        ) from e

    if schema not in ALLOWED_SCHEMAS:
        raise HTTPException(status_code=400, detail="Invalid schema")
    assert_ident(table_name, "table_name")
    lim = min(max(1, limit), 10000)
    so = normalize_sort_order_param(sort_order)

    with db_conn() as conn:
        with conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor) as cur:
            meta = fetch_column_meta(cur, schema, table_name)
            if not meta:
                raise HTTPException(status_code=404, detail="Table not found")

            col_order = _resolve_export_columns(meta, cols)
            name_to_meta = {r["name"]: r for r in meta}

            cur.execute(
                """
                SELECT field_name, COALESCE(menu_label_cn, field_name) AS menu_label_cn
                FROM meta.meta_fields
                WHERE table_name = %s
                """,
                (table_name,),
            )
            label_map = {r["field_name"]: r["menu_label_cn"] for r in cur.fetchall()}
            headers = [label_map.get(c, c) for c in col_order]

            try:
                db_rows_dicts, _ = fetch_table_page(
                    cur,
                    schema,
                    table_name,
                    meta,
                    filters,
                    q,
                    sort_by,
                    so,
                    lim,
                    0,
                    with_total=False,
                )
            except HTTPException:
                raise
            except Exception as e:
                raise HTTPException(status_code=400, detail=str(e)) from e

    wb = Workbook()
    ws = wb.active
    ws.title = "data"[:31]
    ws.append(headers)
    for d in db_rows_dicts:
        ws.append(
            [_export_cell_value(d.get(c), name_to_meta.get(c)) for c in col_order]
        )

    buf = io.BytesIO()
    wb.save(buf)
    body = buf.getvalue()
    fname = f"{table_name}.xlsx"
    # 不用 StreamingResponse(BytesIO)：迭代器可能按行切分二进制，损坏 xlsx
    return Response(
        content=body,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{fname}"'},
    )


@app.post("/table/{schema}/{table_name}")
def table_insert(schema: str, table_name: str, body: Dict[str, Any] = Body(...)):
    """通用插入；设计期全表开放，按可写字段过滤。"""
    if schema not in ALLOWED_SCHEMAS:
        raise HTTPException(status_code=400, detail="Invalid schema")
    assert_ident(schema, "schema")
    assert_ident(table_name, "table_name")
    with db_conn() as conn:
        with conn:
            with conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor) as cur:
                meta = fetch_column_meta(cur, schema, table_name)
                if not meta:
                    raise HTTPException(status_code=404, detail="Table not found")
                allowed = {r["name"]: r for r in writable_columns(meta)}
                cols = []
                vals = []
                for k, v in body.items():
                    if k not in allowed:
                        continue
                    r = allowed[k]
                    if k == "id" and (v is None or v == ""):
                        continue
                    cols.append(k)
                    vals.append(coerce_cell(v, r["udt_name"], r["data_type"], r["is_nullable"]))
                if not cols:
                    raise HTTPException(status_code=400, detail="No writable columns in body")
                qcols = ",".join(f'"{c}"' for c in cols)
                ph = ",".join(["%s"] * len(cols))
                sql = f'INSERT INTO "{schema}"."{table_name}" ({qcols}) VALUES ({ph}) RETURNING *'
                try:
                    cur.execute(sql, vals)
                    row = cur.fetchone()
                except Exception as e:
                    raise HTTPException(status_code=400, detail=str(e)) from e
        return {"row": jsonable_encoder(dict(row))}


@app.patch("/table/{schema}/{table_name}/{row_id}")
def table_update(schema: str, table_name: str, row_id: int, body: Dict[str, Any] = Body(...)):
    if schema not in ALLOWED_SCHEMAS:
        raise HTTPException(status_code=400, detail="Invalid schema")
    assert_ident(schema, "schema")
    assert_ident(table_name, "table_name")
    with db_conn() as conn:
        with conn:
            with conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor) as cur:
                meta = fetch_column_meta(cur, schema, table_name)
                if not meta or not table_has_id(meta):
                    raise HTTPException(status_code=400, detail="Table must have integer id column")
                allowed = {r["name"]: r for r in writable_columns(meta)}
                sets = []
                vals = []
                for k, v in body.items():
                    if k == "id":
                        continue
                    if k not in allowed:
                        continue
                    r = allowed[k]
                    sets.append(f'"{k}" = %s')
                    vals.append(coerce_cell(v, r["udt_name"], r["data_type"], r["is_nullable"]))
                if not sets:
                    raise HTTPException(status_code=400, detail="No updatable fields")
                vals.append(row_id)
                sql = f'UPDATE "{schema}"."{table_name}" SET {", ".join(sets)} WHERE id = %s RETURNING *'
                try:
                    cur.execute(sql, vals)
                    row = cur.fetchone()
                    if not row:
                        raise HTTPException(status_code=404, detail="Row not found")
                except HTTPException:
                    raise
                except Exception as e:
                    raise HTTPException(status_code=400, detail=str(e)) from e
        return {"row": jsonable_encoder(dict(row))}


@app.delete("/table/{schema}/{table_name}/{row_id}")
def table_delete(schema: str, table_name: str, row_id: int):
    if schema not in ALLOWED_SCHEMAS:
        raise HTTPException(status_code=400, detail="Invalid schema")
    assert_ident(schema, "schema")
    assert_ident(table_name, "table_name")
    with db_conn() as conn:
        with conn:
            with conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor) as cur:
                meta = fetch_column_meta(cur, schema, table_name)
                if not meta or not table_has_id(meta):
                    raise HTTPException(status_code=400, detail="Table must have integer id column")
                try:
                    cur.execute(
                        f'DELETE FROM "{schema}"."{table_name}" WHERE id = %s RETURNING *',
                        (row_id,),
                    )
                    row = cur.fetchone()
                    if not row:
                        raise HTTPException(status_code=404, detail="Row not found")
                except HTTPException:
                    raise
                except Exception as e:
                    raise HTTPException(status_code=400, detail=str(e)) from e
        return {"row": jsonable_encoder(dict(row))}

# ── AI Chat ─────────────────────────────────────────────────────────────────
from llm.client import parse_file as parse_llm_file
from llm.tools import (
    DbStreamPendingMarker,
    chat_with_db_outcome,
    chat_with_db_outcome_stream,
    resume_pending_db_chat,
    resume_pending_db_chat_stream,
)
from engine.intent import (
    detect_intent,
    infer_target_table_from_user_text,
    is_business_message,
)
from engine.collecting import COLLECTING_CUD_INTENTS, SLOT_SOURCE_AGENT_PREVIEW, run_free_chat
from engine.workflow_engine import run_confirm_execution
from engine.turns import append_turn, get_recent_turns
from engine.workflow_runner import execute_workflow_for_chat


def _sse_delta_payload(text: str) -> bytes:
    return f"data: {json.dumps({'type': 'delta', 'text': text}, ensure_ascii=False)}\n\n".encode(
        "utf-8"
    )


class ChatMessage(BaseModel):
    role: str
    content: str

class ConfirmRequest(BaseModel):
    frame_id: int
    confirmed: bool
    user_id: Optional[int] = None


class FeedbackRequest(BaseModel):
    session_id: Optional[int] = None
    frame_id: Optional[int] = None
    user_id: Optional[int] = None
    user_comment: str = ""
    feedback_type: str = "bug"
    error_message: Optional[str] = None
    page_url: Optional[str] = None


class ChatRequest(BaseModel):
    model_config = ConfigDict(populate_by_name=True)

    messages: List[ChatMessage]
    # 预留：未来支持图片（base64 列表）
    images: Optional[List[str]] = None
    # 文件解析输入：前端传入 "<filename>|<base64>" 字符串数组（全内存处理，无临时文件）
    files: Optional[List[str]] = None
    # 用户点【确认】：无 pendingToken 时用于认可复述后执行 SQL（遗留）；有 pendingToken 时走二次确认工具
    confirm: Optional[bool] = False
    # 删除/删表二次确认：与 confirm 同时为 true 时恢复执行挂起的工具批次
    pending_token: Optional[str] = Field(default=None, alias="pendingToken")
    # True 时返回 text/event-stream（SSE），正文以增量推送；意图判断仍先完整跑一轮
    stream: Optional[bool] = False
    # True 时本轮对话仅使用云端专家模型（DeepSeek），不自动从本地 fallback
    invite_expert: Optional[bool] = Field(default=False, alias="inviteExpert")
    session_id: Optional[int] = None
    user_id: Optional[int] = None


def _sync_get_or_create_session(user_id: int) -> int:
    """在 ``db_conn()`` 上下文中查找或创建 session，同步实现供 asyncio.to_thread 调用。"""
    with db_conn() as conn:
        conn.autocommit = True
        with conn.cursor() as cur:
            cur.execute(
                """
                SELECT id FROM business.sessions
                WHERE user_id = %s
                  AND is_active = true
                  AND last_active_at > now() - interval '2 hours'
                ORDER BY last_active_at DESC
                LIMIT 1
                """,
                (user_id,),
            )
            row = cur.fetchone()
            if row:
                sid = int(row[0])
                cur.execute(
                    "UPDATE business.sessions SET last_active_at = now() WHERE id = %s",
                    (sid,),
                )
                return sid
            cur.execute(
                """
                INSERT INTO business.sessions (user_id)
                VALUES (%s)
                RETURNING id
                """,
                (user_id,),
            )
            new_row = cur.fetchone()
            if not new_row:
                raise RuntimeError("INSERT business.sessions 未返回 id")
            return int(new_row[0])


def _sync_validate_and_touch_session(user_id: int, session_id: int) -> int:
    """
    客户端显式传入 session_id 时使用：校验归属与有效，并刷新 last_active_at。
    用于 e2e / 前端「新开会话」后指定 id，避免与旧会话合并。
    """
    with db_conn() as conn:
        conn.autocommit = True
        with conn.cursor() as cur:
            cur.execute(
                """
                SELECT id FROM business.sessions
                WHERE id = %s AND user_id = %s
                  AND COALESCE(is_active, true) = true
                  AND deleted_at IS NULL
                """,
                (session_id, user_id),
            )
            row = cur.fetchone()
            if not row:
                raise ValueError("会话不存在、已删除或不属于当前用户")
            sid = int(row[0])
            cur.execute(
                "UPDATE business.sessions SET last_active_at = now() WHERE id = %s",
                (sid,),
            )
            return sid


def _confirm_http_json(
    *,
    content: str,
    wf_intent: str,
    confirm_result: str,
    workflow_context: Optional[Dict[str, Any]] = None,
    error_detail: Optional[str] = None,
) -> Dict[str, Any]:
    """统一 /ai/confirm JSON：成功 / 失败 / 取消 均带 confirm_result，失败带 error_detail。"""
    out: Dict[str, Any] = {
        "content": content,
        "intent": wf_intent,
        "requiresConfirm": False,
        "pendingToken": None,
        "confirm_result": confirm_result,
    }
    if workflow_context is not None:
        out["workflow_context"] = workflow_context
    if error_detail is not None:
        out["error_detail"] = error_detail
    return out


def _sync_ai_confirm(frame_id: int, confirmed: bool, user_id: int) -> Dict[str, Any]:
    """
    POST /ai/confirm：用户确认或取消待写入的 collecting 预览。
    confirmed=True 时执行 workflow_engine.run_confirm_execution（按 frame.workflow_intent
    读 workflow_steps 的 write 步骤 + validation_rules.composite；payload 仍来自 pending_submit）；
    False 时将 frame 置回 collecting，并删除预览 slots（source = agent_preview）。
    """
    with db_conn() as conn:
        conn.autocommit = True
        with conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor) as cur:
            cur.execute(
                """
                SELECT id, user_id, status, context_snapshot, subject, workflow_intent
                FROM business.frames
                WHERE id = %s
                """,
                (int(frame_id),),
            )
            row = cur.fetchone()
            if not row:
                raise HTTPException(status_code=404, detail="frame 不存在")
            if int(row["user_id"]) != int(user_id):
                raise HTTPException(status_code=403, detail="无权操作该任务")

            db_status = row.get("status")
            if (str(db_status or "").strip().lower() == "completed"):
                logger.warning(
                    "confirm: frame id=%s 已为 completed，拒绝重复执行",
                    int(frame_id),
                )
                raise HTTPException(status_code=400, detail="该操作已完成")
            snap_raw = row.get("context_snapshot")
            snap_for_log: Any = snap_raw
            if isinstance(snap_raw, (dict, list)):
                snap_for_log = snap_raw
            elif isinstance(snap_raw, str):
                try:
                    snap_for_log = json.loads(snap_raw)
                except json.JSONDecodeError:
                    snap_for_log = snap_raw
            try:
                snap_json = json.dumps(
                    snap_for_log, ensure_ascii=False, default=str, indent=None
                )
            except TypeError:
                snap_json = str(snap_for_log)
            if len(snap_json) > 8000:
                snap_json = snap_json[:8000] + "…(truncated)"
            logger.info(
                "confirm: 库内 frame id=%s status=%r （待 UI 确认时多为 collecting；"
                "pending_confirm 通常在接口返回的 workflow_context，不一定写入本列）",
                int(frame_id),
                db_status,
            )
            logger.info(
                "confirm: context_snapshot=%s",
                snap_json,
            )
            logger.info(
                "confirm: status==pending_confirm ? %s （当前实现不依赖此列分支，"
                "以 context_snapshot.action / pending_update 为准）",
                (str(db_status).strip().lower() == "pending_confirm"),
            )

            wf_intent = (row.get("workflow_intent") or "").strip() or "create_simple_entity"

            if not confirmed:
                logger.info("confirm: entering cancel branch (confirmed=False)")
                cur.execute(
                    """
                    UPDATE business.frames
                    SET status = 'collecting',
                        updated_at = now(),
                        context_snapshot = (
                            COALESCE(context_snapshot, '{}'::jsonb)
                            - 'pending_update'
                            - 'pending_delete'
                            - 'pending_contract'
                            - 'pending_contract_update'
                            - 'pending_submit'
                            - 'action'
                        )
                    WHERE id = %s AND user_id = %s
                    """,
                    (int(frame_id), int(user_id)),
                )
                cur.execute(
                    """
                    DELETE FROM business.slots
                    WHERE frame_id = %s AND source = %s
                    """,
                    (int(frame_id), SLOT_SOURCE_AGENT_PREVIEW),
                )
                logger.info(
                    "[confirm.取消] frame=%s cleaned_keys=%s",
                    int(frame_id),
                    [
                        "pending_update",
                        "pending_delete",
                        "pending_contract",
                        "pending_contract_update",
                        "pending_submit",
                        "action",
                    ],
                )
                return _confirm_http_json(
                    content="已取消，请重新描述",
                    wf_intent=wf_intent,
                    confirm_result="cancelled",
                )

            logger.info("confirm: entering confirmed=True 分支（V15 workflow_engine）")

            snap = row.get("context_snapshot")
            if isinstance(snap, str):
                try:
                    snap = json.loads(snap)
                except json.JSONDecodeError:
                    snap = {}
            if not isinstance(snap, dict):
                snap = {}

            ps = snap.get("pending_submit")
            if not isinstance(ps, dict):
                logger.warning(
                    "confirm: 缺少 pending_submit，keys=%s",
                    list(snap.keys()),
                )
                raise HTTPException(
                    status_code=400,
                    detail="缺少待确认数据（pending_submit），请重新发起对话",
                )

            logger.info(
                "[confirm.snapshot] frame=%s pending_submit.keys=%s",
                int(frame_id),
                list(ps.keys()),
            )
            _wlist = ps.get("writes") if isinstance(ps.get("writes"), list) else []
            logger.info(
                "[confirm.writes] frame=%s writes_count=%s writes=%s",
                int(frame_id),
                len(_wlist),
                json.dumps(_wlist, ensure_ascii=False, default=str)[:8000],
            )

            ctx_wf: Dict[str, Any] = {
                "frame_id": int(frame_id),
                "user_id": int(user_id),
                "subject": (row.get("subject") or "") if row else "",
            }
            run_confirm_execution(int(frame_id), int(user_id), ctx_wf)
            reply = str(ctx_wf.get("reply") or "")
            if str(ctx_wf.get("status") or "").strip().lower() == "error":
                logger.info(
                    "[confirm.结果] frame=%s success=False detail=%s",
                    int(frame_id),
                    reply[:500],
                )
                return _confirm_http_json(
                    content=reply,
                    wf_intent=wf_intent,
                    confirm_result="error",
                    workflow_context=ctx_wf,
                    error_detail=reply or "执行失败",
                )
            logger.info(
                "[confirm.结果] frame=%s success=True detail=%s",
                int(frame_id),
                reply[:300],
            )
            return _confirm_http_json(
                content=reply,
                wf_intent=wf_intent,
                confirm_result="success",
                workflow_context=ctx_wf,
            )



def _sync_get_collecting_frame_intent(
    user_id: int,
) -> Optional[Tuple[int, str, Optional[str]]]:
    """
    若存在 status=collecting 的 frame，返回 (frame_id, workflow_intent, target_table)，否则 None。
    target_table 来自 context_snapshot，用于与当前用户话推断表比对、避免旧 frame 锁死话题。
    """
    with db_conn() as conn:
        with conn.cursor() as cur:
            cur.execute(
                """
                SELECT id, workflow_intent, context_snapshot
                FROM business.frames
                WHERE user_id = %s AND status = 'collecting'
                ORDER BY updated_at DESC
                LIMIT 1
                """,
                (user_id,),
            )
            row = cur.fetchone()
            if not row:
                return None
            wid = (row[1] or "").strip()
            if not wid:
                return None
            snap = row[2]
            tt: Optional[str] = None
            if snap is not None:
                if isinstance(snap, str):
                    try:
                        snap = json.loads(snap)
                    except json.JSONDecodeError:
                        snap = {}
                if isinstance(snap, dict):
                    raw_tt = snap.get("target_table")
                    if isinstance(raw_tt, str) and raw_tt.strip():
                        tt = raw_tt.strip()
            return (int(row[0]), wid, tt)


def _sync_supersede_collecting_frame(frame_id: int, user_id: int) -> None:
    """
    结束进行中的 collecting frame（意图或目标表变化，不再参与后续路由）。
    status 必须为业务.check 允许的值之一：collecting / completed / cancelled（不可用 superseded）。
    """
    with db_conn() as conn:
        conn.autocommit = True
        with conn.cursor() as cur:
            cur.execute(
                """
                UPDATE business.frames
                SET status = 'cancelled',
                    subject = COALESCE(subject, '') || ' [话题已切换]',
                    updated_at = now()
                WHERE id = %s AND user_id = %s AND status = 'collecting'
                """,
                (int(frame_id), int(user_id)),
            )


def _workflow_messages_history(history: List[Dict[str, Any]], limit: int = 10) -> List[Dict[str, Any]]:
    """取最近 limit 条对话，供 collecting 填槽。"""
    if not history:
        return []
    return history[-limit:] if len(history) > limit else list(history)


async def get_or_create_session(user_id: int) -> int:
    """
    找到该用户最近一个活跃 session（last_active_at 在 2 小时内）则复用并刷新；
    否则新建。使用 API 连接池 ``db_conn()``。
    """
    return await asyncio.to_thread(_sync_get_or_create_session, user_id)


@app.post("/ai/confirm")
async def ai_confirm(body: ConfirmRequest):
    """collecting 待确认写入：确认则执行 INSERT，取消则保持 collecting。"""
    frame_id = body.frame_id
    confirmed = body.confirmed
    uid = body.user_id or 2
    logger.info(
        "[confirm.入口] frame=%s confirmed=%s user=%s",
        frame_id,
        confirmed,
        uid,
    )
    return await asyncio.to_thread(_sync_ai_confirm, frame_id, confirmed, uid)


def _sync_ai_feedback(body: FeedbackRequest) -> Dict[str, Any]:
    """写入 business.feedback，附最近 turns 与 frame 快照。"""
    uid = body.user_id or 2
    recent_turns: List[Any] = []
    frame_snapshot: Any = {}
    intent: Optional[str] = None

    with db_conn() as conn:
        conn.autocommit = True
        if body.session_id is not None:
            with conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor) as cur:
                cur.execute(
                    """
                    SELECT role, content, created_at
                    FROM business.turns
                    WHERE session_id = %s
                    ORDER BY created_at DESC
                    LIMIT 10
                    """,
                    (int(body.session_id),),
                )
                rows = cur.fetchall()
                recent_turns = [dict(r) for r in rows]
                recent_turns.reverse()

        if body.frame_id is not None:
            with conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor) as cur:
                cur.execute(
                    """
                    SELECT workflow_intent, status, context_snapshot
                    FROM business.frames
                    WHERE id = %s
                    """,
                    (int(body.frame_id),),
                )
                row = cur.fetchone()
                if row:
                    intent = row.get("workflow_intent")
                    if intent is not None:
                        intent = str(intent).strip() or None
                    raw_snap = row.get("context_snapshot")
                    if isinstance(raw_snap, dict):
                        frame_snapshot = raw_snap
                    elif isinstance(raw_snap, str):
                        try:
                            frame_snapshot = (
                                json.loads(raw_snap) if raw_snap.strip() else {}
                            )
                        except json.JSONDecodeError:
                            frame_snapshot = {}
                    else:
                        frame_snapshot = {}

        with conn.cursor() as cur:
            cur.execute(
                """
                INSERT INTO business.feedback
                (session_id, frame_id, user_id, user_comment, feedback_type,
                 recent_turns, frame_snapshot, intent, error_message, page_url)
                VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
                RETURNING id
                """,
                (
                    body.session_id,
                    body.frame_id,
                    uid,
                    body.user_comment,
                    body.feedback_type,
                    Json(recent_turns),
                    Json(frame_snapshot if isinstance(frame_snapshot, dict) else {}),
                    intent,
                    body.error_message,
                    body.page_url,
                ),
            )
            row_id = cur.fetchone()
            if not row_id:
                raise RuntimeError("INSERT business.feedback 未返回 id")
            feedback_id = int(row_id[0])

    logger.info(
        "[feedback.写入] id=%s session=%s frame=%s type=%s",
        feedback_id,
        body.session_id,
        body.frame_id,
        body.feedback_type,
    )
    return {"success": True, "feedback_id": feedback_id}


@app.post("/ai/feedback")
async def ai_feedback(body: FeedbackRequest):
    logger.info(
        "[feedback.入口] session=%s frame=%s type=%s",
        body.session_id,
        body.frame_id,
        body.feedback_type,
    )
    try:
        return await asyncio.to_thread(_sync_ai_feedback, body)
    except Exception as e:
        logger.exception("[feedback] 写入失败: %s", e)
        raise HTTPException(status_code=500, detail=str(e)) from e


def get_llm_db_config() -> Dict[str, Any]:
    """供 LLM 工具调用的库连接（与 setup_llm_user 中助手账号一致）。"""
    port_raw = os.environ.get("PGPORT", "5432")
    try:
        port = int(port_raw)
    except ValueError:
        port = 5432
    return {
        "host": os.environ.get("PGHOST", "localhost"),
        "database": os.environ.get("PGDATABASE", "aierp"),
        "user": os.environ.get("PGUSER", "llm_assistant"),
        "password": os.environ.get("PGPASSWORD", "llm_safe_2024"),
        "port": port,
    }


def _db_chat_system_prompt(base_system: str) -> str:
    return (
        base_system
        + "\n\n你与用户的对话始终在「可调数据库工具」的环境下进行，无需事先区分闲聊或查库；"
        "若问题与数据库、业务数据无关，可直接用文字回答，不必调用工具。"
        "一旦涉及查看或变更库中数据、执行 SQL、统计业务表等，必须调用 query_db / execute_db，"
        "严禁只用 Markdown 代码块展示 SQL 而不调用工具，勿编造「已执行」。"
        "query_db 仅允许单条 SELECT；execute_db 允许 INSERT、UPDATE、DELETE、CREATE、DROP。"
        "根据工具返回的 JSON（如 rows_affected、error）向用户说明结果。"
        "若用户要删除或删除助手本人所建表中的数据，系统会要求用户在界面确认后再执行。"
        "回答时用简洁中文，必要时使用 Markdown（代码块仅作补充说明，不能代替工具）。"
        "\n\n遇到工具执行报错时：\n"
        "1. 仔细阅读 error 和 hint 字段\n"
        "2. 先用 query_db 查询相关表结构或现有数据，获取足够上下文\n"
        "3. 根据查询结果调整后重新执行，不要直接放弃或只回复错误信息\n"
        "4. 如果多次尝试仍失败，用业务语言告知用户原因，不要暴露原始SQL和技术细节"
    )


def _db_chat_after_operation_user_confirm_prompt(base_system: str) -> str:
    """用户已点【确认】认可复述后，驱动模型实际执行 SQL。"""
    return (
        _db_chat_system_prompt(base_system)
        + "\n\n【重要】用户已在界面点击「确认」，认可你刚才复述的数据库操作意图。"
        "请立即使用 query_db / execute_db 完成对应数据库变更，并简要汇报结果。"
        "不要再次向用户索要确认；若系统因「删除助手自建对象」返回需二次确认，"
        "按工具返回的说明在界面完成确认即可。"
    )

@app.post("/ai/chat")
async def ai_chat(req: ChatRequest):
    # —— 调试：完整请求体（含 pendingToken 别名解析结果）——
    try:
        _req_dump = req.model_dump(mode="json", by_alias=True)
    except Exception:
        _req_dump = jsonable_encoder(req)
    print(
        "[ai/chat] 请求体:",
        json.dumps(_req_dump, ensure_ascii=False, default=str),
        flush=True,
    )

    # 构建基础系统提示：注入数据库上下文
    base_system = """你是 AI·ERP 系统的智能助手。你熟悉该系统的数据库结构：
- meta schema：仅保留元数据目录表 meta_tables、meta_fields（表/字段说明、导航配置等）
- business schema：正式业务与主数据（含 ai_rules、employees、products、合同与库存等表及视图），可查询或按权限更新，但不要在 business 下新建表
- ai_assistant schema：助手专属工作区；凡由你新建的数据表、临时表、实验表必须建在此 schema 下（名称一律用小写 ai_assistant，勿用 business）

请用简洁、专业的中文回答问题。涉及数据或代码时，使用 Markdown 格式。"""

    history_messages = [{"role": m.role, "content": m.content} for m in req.messages]

    # 1) 先解析文件（如果有）
    file_text_parts: List[str] = []
    file_images_b64: List[str] = []
    file_summaries: List[str] = []

    if req.files:
        for file_payload in req.files:
            parsed = await asyncio.to_thread(parse_llm_file, file_payload)
            p_type = parsed.get("type")
            if p_type == "images":
                file_images_b64.extend(parsed.get("content", []) or [])
            elif p_type == "text":
                file_text_parts.append(parsed.get("content", "") or "")
            if parsed.get("summary"):
                file_summaries.append(str(parsed["summary"]))

    file_text_prompt = ""
    if file_text_parts:
        file_text_prompt = (
            "用户上传了附件，下面是附件解析结果（供模型理解）：\n"
            + "\n\n".join(file_text_parts)
        )

    # 把文字解析结果拼进 prompt（挂到最后一条 user 消息上）
    if file_text_prompt:
        for m in reversed(history_messages):
            if m.get("role") == "user":
                m["content"] = m.get("content", "") + "\n\n" + file_text_prompt
                break

    # 1b) 意图路由：检查用户输入是否匹配 workflow
    # 只对普通新消息做意图识别（confirm 分支、pending 分支跳过）
    if not req.confirm and not req.pending_token:
        last_user_text = ""
        for m in reversed(history_messages):
            if m.get("role") == "user":
                last_user_text = m.get("content", "")
                break

        user_id = req.user_id or 2  # TODO: 登录模块完成后替换
        collecting_frame: Optional[Tuple[int, str, Optional[str]]] = None
        detected_intent: Optional[str] = None
        matched_intent: Optional[str] = None
        session_id: Optional[int] = None
        intent_sidecar: Dict[str, Any] = {}

        if last_user_text.strip():
            try:
                collecting_frame = await asyncio.to_thread(
                    _sync_get_collecting_frame_intent, user_id
                )
            except Exception as _cfe:
                print(f"[ai/chat] 查询 collecting frame 异常: {_cfe}", flush=True)

            if req.session_id is not None:
                try:
                    sid_raw = int(req.session_id)
                except (TypeError, ValueError) as _e:
                    raise HTTPException(
                        status_code=400, detail="session_id 须为整数"
                    ) from _e
                try:
                    session_id = await asyncio.to_thread(
                        _sync_validate_and_touch_session, user_id, sid_raw
                    )
                except ValueError as _e:
                    raise HTTPException(status_code=400, detail=str(_e)) from _e
            else:
                session_id = await get_or_create_session(user_id)
            intent_sidecar = {"session_id": session_id}

            try:
                detected_intent = await asyncio.to_thread(
                    detect_intent,
                    last_user_text,
                    bool(req.invite_expert),
                    intent_sidecar,
                )
            except Exception as _ie:
                detected_intent = None
                print(
                    f"[ai/chat] 意图识别异常（忽略，走兜底）: {_ie}",
                    flush=True,
                )

            if collecting_frame and detected_intent:
                _cf_fid, _cf_wi, _cf_tt = collecting_frame
                _wi = _cf_wi.strip()
                _di = detected_intent.strip()
                # 本轮识别为查主数据 / free_chat 且 frame 为 CUD：走轻量或 free_chat 路径，不 supersede、不改 frame
                _skip_supersede_for_free_chat_pivot = _wi in COLLECTING_CUD_INTENTS and _di in (
                    "query_master_data",
                    "free_chat",
                )
                # 合同收集中，用户补充公司名等易被误判为「建客户 / 查主数据」，不 supersede
                _skip_supersede_contract_collecting = _wi in (
                    "create_contract",
                    "update_contract",
                ) and _di in ("create_simple_entity", "query_master_data")
                if (
                    _di != _wi
                    and not _skip_supersede_contract_collecting
                    and not _skip_supersede_for_free_chat_pivot
                ):
                    try:
                        await asyncio.to_thread(
                            _sync_supersede_collecting_frame, _cf_fid, user_id
                        )
                    except Exception as _se:
                        print(
                            f"[ai/chat] supersede（意图不一致）失败: {_se}",
                            flush=True,
                        )
                    else:
                        logger.info(
                            "[chat.supersede] old_frame=%s old_intent=%s new_intent=%s reason=%s",
                            _cf_fid,
                            _cf_wi,
                            _di,
                            "intent不一致",
                        )
                        print(
                            f"[ai/chat] detect_intent={detected_intent!r} 与 collecting frame "
                            f"workflow_intent={_cf_wi!r} 不一致，已 supersede frame id={_cf_fid}",
                            flush=True,
                        )
                        collecting_frame = None

            # frame 为 CUD 且本轮识别为 None 或 query_master_data：走轻量路径，不因 target_table 推断 supersede frame
            _skip_infer_supersede_for_free_chat = False
            if collecting_frame:
                _fw_tmp = collecting_frame[1].strip()
                if _fw_tmp in COLLECTING_CUD_INTENTS:
                    _det_tmp = (detected_intent or "").strip()
                    if not _det_tmp or _det_tmp in (
                        "query_master_data",
                        "free_chat",
                    ):
                        _skip_infer_supersede_for_free_chat = True

            if (
                collecting_frame
                and last_user_text.strip()
                and not _skip_infer_supersede_for_free_chat
            ):
                _cf_fid, _cf_wi, frame_tt = collecting_frame
                _skip_infer_supersede = _cf_wi.strip() in (
                    "create_contract",
                    "update_contract",
                    "delete_contract",
                )
                try:
                    inferred_tt = await asyncio.to_thread(
                        infer_target_table_from_user_text, last_user_text
                    )
                except Exception as _inf_e:
                    inferred_tt = None
                    print(
                        f"[ai/chat] 推断 target_table 异常（忽略）: {_inf_e}",
                        flush=True,
                    )
                if (
                    not _skip_infer_supersede
                    and inferred_tt
                    and frame_tt
                    and inferred_tt.strip() != frame_tt.strip()
                ):
                    try:
                        await asyncio.to_thread(
                            _sync_supersede_collecting_frame, _cf_fid, user_id
                        )
                    except Exception as _se:
                        print(
                            f"[ai/chat] 结束旧 collecting frame 失败: {_se}",
                            flush=True,
                        )
                    else:
                        logger.info(
                            "[chat.supersede] old_frame=%s old_intent=%s new_intent=%s reason=%s",
                            _cf_fid,
                            _cf_wi,
                            inferred_tt,
                            "target_table不一致",
                        )
                        print(
                            f"[ai/chat] 用户话推断表={inferred_tt!r} 与 frame 目标表={frame_tt!r} 不一致，"
                            f"已结束 frame id={_cf_fid}，重新走意图识别",
                            flush=True,
                        )
                        collecting_frame = None

            if collecting_frame:
                _, frame_wi, _ = collecting_frame
                fw = frame_wi.strip()
                di = (detected_intent or "").strip()
                # 与 supersede 豁免一致：合同收集中误判为建客户/查主数据时，仍走 frame 意图
                if fw in ("create_contract", "update_contract") and di in (
                    "create_simple_entity",
                    "query_master_data",
                ):
                    matched_intent = fw
                elif di:
                    matched_intent = di
                else:
                    matched_intent = fw
            else:
                matched_intent = (
                    detected_intent.strip() if detected_intent else None
                )

            if collecting_frame:
                print(
                    f"[ai/chat] detect_intent={detected_intent!r}，复用 collecting frame "
                    f"id={collecting_frame[0]}，matched_intent={matched_intent!r}",
                    flush=True,
                )
            else:
                print(
                    f"[ai/chat] detect_intent={detected_intent!r}，matched_intent={matched_intent!r} "
                    f"（无 collecting frame 或已 supersede）",
                    flush=True,
                )

        _d_gate = (detected_intent or "").strip()
        if last_user_text.strip() and (
            collecting_frame is not None
            or is_business_message(last_user_text)
            or not _d_gate
            or _d_gate == "query_master_data"
            or _d_gate == "free_chat"
        ):
            if session_id is None:
                session_id = await get_or_create_session(user_id)
            intent_sidecar = {"session_id": session_id}
            # 分流仅以本轮 detect_intent 为准，不用 matched_intent（后者可能来自旧 frame）
            dispatch_intent = (detected_intent or "").strip() or None
            workflow_ctx = {
                "user_id": user_id,
                "session_id": session_id,
                "role": getattr(req, "role", None),
                "user_input": last_user_text,
                "messages_history": _workflow_messages_history(
                    history_messages, limit=10
                ),
            }
            if not collecting_frame:
                if intent_sidecar.get("target_table"):
                    workflow_ctx["target_table"] = str(
                        intent_sidecar["target_table"]
                    ).strip()
                if intent_sidecar.get("intent_reason"):
                    workflow_ctx["intent_reason"] = str(
                        intent_sidecar["intent_reason"]
                    ).strip()
            # 如果有文件解析内容，也带入
            if file_text_parts:
                workflow_ctx["uploaded_file_text"] = "\n\n".join(file_text_parts)

            await append_turn(session_id, None, "user", last_user_text)

            logger.info(
                "[chat.入口] session=%s detected=%s dispatch=%s frame=%s matched_legacy=%s",
                session_id,
                detected_intent,
                dispatch_intent,
                collecting_frame[0] if collecting_frame else None,
                matched_intent,
            )
            try:
                if dispatch_intent in COLLECTING_CUD_INTENTS:
                    result = await asyncio.to_thread(
                        execute_workflow_for_chat, dispatch_intent, workflow_ctx
                    )
                elif dispatch_intent is None or dispatch_intent == "free_chat":
                    result = await asyncio.to_thread(
                        execute_workflow_for_chat, "free_chat", workflow_ctx
                    )
                elif dispatch_intent == "query_master_data":
                    workflow_ctx["free_chat_intent_label"] = "query_master_data"
                    result = await asyncio.to_thread(run_free_chat, workflow_ctx)
                else:
                    result = await asyncio.to_thread(
                        execute_workflow_for_chat, dispatch_intent, workflow_ctx
                    )
            except Exception as _we:
                _fail_intent = dispatch_intent or "未匹配"
                result = {
                    "content": f"流程执行失败：{_we}",
                    "intent": _fail_intent,
                    "requiresConfirm": False,
                    "pendingToken": None,
                }

            reply_content = str(
                result.get("reply") or result.get("content") or ""
            )
            frame_id = workflow_ctx.get("frame_id")
            await append_turn(session_id, frame_id, "assistant", reply_content)

            result.setdefault("fileSummaries", file_summaries)
            result["session_id"] = session_id
            wc = result.get("workflow_context")
            if isinstance(wc, dict) and wc.get("status") == "pending_confirm":
                result["confirm_required"] = True
                result["preview"] = wc.get("preview")
                result["explanation"] = wc.get("explanation")
                if wc.get("frame_id") is not None:
                    result["frame_id"] = wc.get("frame_id")
            if result.get("frame_id") is None:
                if isinstance(wc, dict) and wc.get("frame_id") is not None:
                    result["frame_id"] = wc.get("frame_id")
                elif workflow_ctx.get("frame_id") is not None:
                    result["frame_id"] = workflow_ctx.get("frame_id")
            _wc = result.get("workflow_context")
            _st = (
                (_wc.get("status") if isinstance(_wc, dict) else None)
                or result.get("status")
            )
            logger.info(
                "[chat.出口] frame=%s status=%s has_confirm=%s reply_len=%s",
                result.get("frame_id") or workflow_ctx.get("frame_id"),
                _st,
                bool(result.get("confirm_required")),
                len(reply_content),
            )
            if bool(req.stream) and (
                dispatch_intent == "query_master_data"
                or dispatch_intent is None
                or dispatch_intent == "free_chat"
            ):
                _meta_intent = str(
                    result.get("intent")
                    or (
                        "query_master_data"
                        if dispatch_intent == "query_master_data"
                        else "未匹配"
                    )
                )

                async def sse_free_chat_path():
                    try:
                        yield f"data: {json.dumps({'type': 'start'})}\n\n".encode(
                            "utf-8"
                        )
                        yield f"data: {json.dumps({'type': 'meta', 'intent': _meta_intent, 'requiresConfirm': False, 'pendingToken': None, 'fileSummaries': file_summaries, 'session_id': session_id}, ensure_ascii=False)}\n\n".encode(
                            "utf-8"
                        )
                        yield _sse_delta_payload(reply_content)
                        yield f"data: {json.dumps({'type': 'done'}, ensure_ascii=False)}\n\n".encode(
                            "utf-8"
                        )
                    except Exception as e:
                        err = {"type": "error", "detail": str(e)}
                        yield f"data: {json.dumps(err, ensure_ascii=False)}\n\n".encode(
                            "utf-8"
                        )
                        yield f"data: {json.dumps({'type': 'done'}, ensure_ascii=False)}\n\n".encode(
                            "utf-8"
                        )

                return StreamingResponse(
                    sse_free_chat_path(),
                    media_type="text/event-stream",
                    headers={
                        "Cache-Control": "no-cache",
                        "Connection": "keep-alive",
                        "X-Accel-Buffering": "no",
                    },
                )

            return result

    # 2a) 数据库工具：删除/删表二次确认（pending_token 由上次查询对话返回）
    print(
        "[ai/chat] 分支预判: confirm=%r pending_token=%r stream=%r"
        % (req.confirm, req.pending_token, req.stream),
        flush=True,
    )
    if bool(req.confirm) and req.pending_token:
        print(
            "[ai/chat] 进入 2a pending_token 分支, token=%r" % (req.pending_token,),
            flush=True,
        )
        if bool(req.stream):

            async def sse_resume_pending():
                try:
                    yield f"data: {json.dumps({'type': 'start'})}\n\n".encode("utf-8")
                    try:
                        meta_emitted = False
                        async for item in resume_pending_db_chat_stream(
                            req.pending_token
                        ):
                            if isinstance(item, DbStreamPendingMarker):
                                if not meta_emitted:
                                    yield f"data: {json.dumps({'type': 'meta', 'intent': '查询', 'requiresConfirm': True, 'pendingToken': item.token, 'fileSummaries': file_summaries}, ensure_ascii=False)}\n\n".encode(
                                        "utf-8"
                                    )
                                    meta_emitted = True
                                continue
                            if not meta_emitted:
                                yield f"data: {json.dumps({'type': 'meta', 'intent': '查询', 'requiresConfirm': False, 'pendingToken': None, 'fileSummaries': file_summaries}, ensure_ascii=False)}\n\n".encode(
                                    "utf-8"
                                )
                                meta_emitted = True
                            if item:
                                yield _sse_delta_payload(item)
                    except (RuntimeError, ValueError) as e:
                        err = {"type": "error", "detail": str(e)}
                        yield f"data: {json.dumps(err, ensure_ascii=False)}\n\n".encode(
                            "utf-8"
                        )
                        yield f"data: {json.dumps({'type': 'done'}, ensure_ascii=False)}\n\n".encode(
                            "utf-8"
                        )
                        return
                    if not meta_emitted:
                        yield f"data: {json.dumps({'type': 'meta', 'intent': '查询', 'requiresConfirm': False, 'pendingToken': None, 'fileSummaries': file_summaries}, ensure_ascii=False)}\n\n".encode(
                            "utf-8"
                        )
                    yield f"data: {json.dumps({'type': 'done'}, ensure_ascii=False)}\n\n".encode(
                        "utf-8"
                    )
                except Exception as e:
                    err = {"type": "error", "detail": str(e)}
                    yield f"data: {json.dumps(err, ensure_ascii=False)}\n\n".encode("utf-8")
                    yield f"data: {json.dumps({'type': 'done'}, ensure_ascii=False)}\n\n".encode(
                        "utf-8"
                    )

            return StreamingResponse(
                sse_resume_pending(),
                media_type="text/event-stream",
                headers={
                    "Cache-Control": "no-cache",
                    "Connection": "keep-alive",
                    "X-Accel-Buffering": "no",
                },
            )

        try:
            content = await asyncio.to_thread(resume_pending_db_chat, req.pending_token)
        except (RuntimeError, ValueError) as e:
            raise HTTPException(status_code=502, detail=str(e)) from e
        return {
            "content": content,
            "intent": "查询",
            "requiresConfirm": False,
            "pendingToken": None,
            "fileSummaries": file_summaries,
        }

    # 2b) 用户已点【确认】（含认可 DDL 复述）：用数据库工具实际执行
    if bool(req.confirm):
        db_messages = [
            {
                "role": "system",
                "content": _db_chat_after_operation_user_confirm_prompt(base_system),
            }
        ] + history_messages

        if bool(req.stream):

            async def sse_operation_confirm():
                try:
                    yield f"data: {json.dumps({'type': 'start'})}\n\n".encode("utf-8")
                    try:
                        meta_emitted = False
                        async for item in chat_with_db_outcome_stream(
                            db_messages,
                            get_llm_db_config(),
                            use_expert=bool(req.invite_expert),
                        ):
                            if isinstance(item, DbStreamPendingMarker):
                                if not meta_emitted:
                                    yield f"data: {json.dumps({'type': 'meta', 'intent': '查询', 'requiresConfirm': True, 'pendingToken': item.token, 'fileSummaries': file_summaries}, ensure_ascii=False)}\n\n".encode(
                                        "utf-8"
                                    )
                                    meta_emitted = True
                                continue
                            if not meta_emitted:
                                yield f"data: {json.dumps({'type': 'meta', 'intent': '查询', 'requiresConfirm': False, 'pendingToken': None, 'fileSummaries': file_summaries}, ensure_ascii=False)}\n\n".encode(
                                    "utf-8"
                                )
                                meta_emitted = True
                            if item:
                                yield _sse_delta_payload(item)
                    except (RuntimeError, ValueError) as e:
                        err = {"type": "error", "detail": str(e)}
                        yield f"data: {json.dumps(err, ensure_ascii=False)}\n\n".encode(
                            "utf-8"
                        )
                        yield f"data: {json.dumps({'type': 'done'}, ensure_ascii=False)}\n\n".encode(
                            "utf-8"
                        )
                        return
                    if not meta_emitted:
                        yield f"data: {json.dumps({'type': 'meta', 'intent': '查询', 'requiresConfirm': False, 'pendingToken': None, 'fileSummaries': file_summaries}, ensure_ascii=False)}\n\n".encode(
                            "utf-8"
                        )
                    yield f"data: {json.dumps({'type': 'done'}, ensure_ascii=False)}\n\n".encode(
                        "utf-8"
                    )
                except Exception as e:
                    err = {"type": "error", "detail": str(e)}
                    yield f"data: {json.dumps(err, ensure_ascii=False)}\n\n".encode("utf-8")
                    yield f"data: {json.dumps({'type': 'done'}, ensure_ascii=False)}\n\n".encode(
                        "utf-8"
                    )

            return StreamingResponse(
                sse_operation_confirm(),
                media_type="text/event-stream",
                headers={
                    "Cache-Control": "no-cache",
                    "Connection": "keep-alive",
                    "X-Accel-Buffering": "no",
                },
            )

        try:
            outcome = await asyncio.to_thread(
                chat_with_db_outcome,
                db_messages,
                get_llm_db_config(),
                use_expert=bool(req.invite_expert),
            )
        except (RuntimeError, ValueError) as e:
            raise HTTPException(status_code=502, detail=str(e)) from e
        if outcome.kind == "pending":
            return {
                "content": outcome.explanation or "",
                "intent": "查询",
                "requiresConfirm": True,
                "pendingToken": outcome.token,
                "fileSummaries": file_summaries,
            }
        return {
            "content": outcome.text or "",
            "intent": "查询",
            "requiresConfirm": False,
            "pendingToken": None,
            "fileSummaries": file_summaries,
        }

    # 4a) 流式：统一走数据库工具链（tool_choice=auto，由模型决定是否调工具）
    if bool(req.stream):
        _ = file_images_b64

        async def sse_bytes():
            try:
                yield f"data: {json.dumps({'type': 'start'})}\n\n".encode("utf-8")
                db_messages = [
                    {"role": "system", "content": _db_chat_system_prompt(base_system)}
                ] + history_messages
                try:
                    meta_emitted = False
                    pending_token_seen: Optional[str] = None
                    async for item in chat_with_db_outcome_stream(
                        db_messages,
                        get_llm_db_config(),
                        use_expert=bool(req.invite_expert),
                    ):
                        if isinstance(item, DbStreamPendingMarker):
                            pending_token_seen = item.token
                            if not meta_emitted:
                                yield f"data: {json.dumps({'type': 'meta', 'intent': '查询', 'requiresConfirm': True, 'pendingToken': item.token, 'fileSummaries': file_summaries}, ensure_ascii=False)}\n\n".encode(
                                    "utf-8"
                                )
                                meta_emitted = True
                            continue
                        if not meta_emitted:
                            yield f"data: {json.dumps({'type': 'meta', 'intent': '查询', 'requiresConfirm': False, 'pendingToken': None, 'fileSummaries': file_summaries}, ensure_ascii=False)}\n\n".encode(
                                "utf-8"
                            )
                            meta_emitted = True
                        if item:
                            yield _sse_delta_payload(item)
                    print(
                        "[ai/chat SSE] 工具链流式结束 pendingToken=%r" % (pending_token_seen,),
                        flush=True,
                    )
                except (RuntimeError, ValueError) as e:
                    err = {"type": "error", "detail": str(e)}
                    yield f"data: {json.dumps(err, ensure_ascii=False)}\n\n".encode("utf-8")
                    yield f"data: {json.dumps({'type': 'done'}, ensure_ascii=False)}\n\n".encode("utf-8")
                    return
                if not meta_emitted:
                    yield f"data: {json.dumps({'type': 'meta', 'intent': '查询', 'requiresConfirm': False, 'pendingToken': None, 'fileSummaries': file_summaries}, ensure_ascii=False)}\n\n".encode(
                        "utf-8"
                    )
                yield f"data: {json.dumps({'type': 'done'}, ensure_ascii=False)}\n\n".encode("utf-8")
            except Exception as e:
                err = {"type": "error", "detail": str(e)}
                yield f"data: {json.dumps(err, ensure_ascii=False)}\n\n".encode("utf-8")
                yield f"data: {json.dumps({'type': 'done'}, ensure_ascii=False)}\n\n".encode("utf-8")

        return StreamingResponse(
            sse_bytes(),
            media_type="text/event-stream",
            headers={
                "Cache-Control": "no-cache",
                "Connection": "keep-alive",
                "X-Accel-Buffering": "no",
            },
        )

    # 3) 非流式：统一走数据库工具链
    _ = file_images_b64
    db_messages = [
        {"role": "system", "content": _db_chat_system_prompt(base_system)}
    ] + history_messages
    try:
        outcome = await asyncio.to_thread(
            chat_with_db_outcome,
            db_messages,
            get_llm_db_config(),
            use_expert=bool(req.invite_expert),
        )
    except (RuntimeError, ValueError) as e:
        raise HTTPException(status_code=502, detail=str(e)) from e
    if outcome.kind == "pending":
        return {
            "content": outcome.explanation or "",
            "intent": "查询",
            "requiresConfirm": True,
            "pendingToken": outcome.token,
            "fileSummaries": file_summaries,
        }
    return {
        "content": outcome.text or "",
        "intent": "查询",
        "requiresConfirm": False,
        "pendingToken": None,
        "fileSummaries": file_summaries,
    }
